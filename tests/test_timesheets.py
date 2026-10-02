"""The Timesheets agent in a real worker. Lina's projects at Sidra
Office Supplies, and her calendar for the week of 7 September 2026 as
the calendar agents list it (Dubai time).
"""

import asyncio
import io

import pytest
from openpyxl import load_workbook

from ai_runtime.execution.executor import FunctionExecutor
from sim.resources import InMemoryResourceProvider

TZ = "+04:00"
WEEK = "2026-W37"


def run(awaitable):
    return asyncio.run(awaitable)


def make():
    provider = InMemoryResourceProvider()
    return FunctionExecutor(provider=provider), provider


def invoke(agents, ex, name, inputs, chat_level=1):
    return run(ex.invoke(agents["timesheets"], name, inputs, chat_level=chat_level))


def ok(agents, ex, name, inputs, chat_level=1):
    result, status = invoke(agents, ex, name, inputs, chat_level)
    assert status == "success", result
    return result


def refused(agents, ex, name, inputs, chat_level=1):
    result, status = invoke(agents, ex, name, inputs, chat_level)
    assert status == "error", result
    return result["error"]


def projects(agents, ex):
    ok(agents, ex, "timesheets.projects.add", {
        "code": "show", "name": "Harbourline showroom fit-out", "client": "Harbourline",
        "budget_hours": 60})
    ok(agents, ex, "timesheets.projects.add", {
        "code": "MOVE", "name": "Riverside office move", "billable": False})
    ok(agents, ex, "timesheets.projects.add", {
        "code": "ADMIN", "name": "Admin and internal", "billable": False})
    for rule in ({"match": "Harbourline", "project_code": "SHOW"},
                 {"match": "harbourline.example", "field": "attendee", "project_code": "SHOW"},
                 {"match": "stand-up", "project_code": "ADMIN"},
                 {"match": "Riverside", "project_code": "MOVE"}):
        ok(agents, ex, "timesheets.projects.add_rule", rule)


def meeting(event_id, title, day, start, end, **extra):
    return {"event_id": event_id, "summary": title,
            "start": f"2026-09-{day:02d}T{start}:00{TZ}", "end": f"2026-09-{day:02d}T{end}:00{TZ}",
            "all_day": False, "attendees": "", "location": "", "status": "confirmed", **extra}


CALENDAR = [
    meeting("ev1", "Ops stand-up", 7, "09:00", "09:30"),
    meeting("ev2", "Showroom layout", 7, "10:00", "11:50", attendees="dana@harbourline.example"),
    meeting("ev3", "Riverside movers — quotes", 8, "14:00", "15:00"),
    meeting("ev4", "Lunch with Omar", 8, "12:00", "13:00"),
    meeting("ev5", "Harbourline stand-up", 9, "09:00", "09:15"),
    {"event_id": "ev6", "summary": "Public holiday", "start": "2026-09-10",
     "end": "2026-09-11", "all_day": True},
    meeting("ev7", "Board call", 10, "16:00", "17:00", status="cancelled"),
]


def log(agents, ex, code, day, hours, note=""):
    return ok(agents, ex, "timesheets.entries.log", {
        "project_code": code, "date": f"2026-09-{day:02d}", "hours": hours, "note": note})


class TestThePackage:
    def test_the_worker_verifies_every_declared_function(self, agents):
        from ai_runtime.agents.worker_handle import WorkerHandle

        agent = agents["timesheets"]
        assert WorkerHandle.probe(agent.environment.python, agent.folder,
                                  agent.manifest.document) == []

    def test_it_keeps_to_the_platform(self, agents):
        """Nothing here leaves the platform; the workbook is the one
        sandboxed change, as it is for Expenses."""
        levels = {f"{t['id']}.{f['id']}": f["permission_level"]
                  for t in agents["timesheets"].manifest.document["tools"]
                  for f in t["functions"]}
        assert {name for name, level in levels.items() if level >= 2} == {"weeks.workbook"}
        assert levels["weeks.workbook"] == 2


class TestProjects:
    def test_a_code_is_one_project(self, agents):
        ex, _ = make()
        projects(agents, ex)
        assert "already" in refused(agents, ex, "timesheets.projects.add",
                                    {"code": "SHOW", "name": "Another"})
        listed = ok(agents, ex, "timesheets.projects.list", {})["projects"]
        show = next(p for p in listed if p["code"] == "SHOW")
        assert (show["billable"], show["budget_hours"], show["logged_hours"]) == (True, 60.0, 0.0)
        assert next(p for p in listed if p["code"] == "MOVE")["billable"] is False

    def test_a_rule_needs_an_active_project(self, agents):
        ex, _ = make()
        projects(agents, ex)
        assert "NOPE" in refused(agents, ex, "timesheets.projects.add_rule",
                                 {"match": "x", "project_code": "NOPE"})
        ok(agents, ex, "timesheets.projects.update", {"code": "MOVE", "status": "archived"})
        assert "MOVE" in refused(agents, ex, "timesheets.projects.add_rule",
                                 {"match": "lease", "project_code": "MOVE"})
        codes = [p["code"] for p in ok(agents, ex, "timesheets.projects.list", {})["projects"]]
        assert codes == ["ADMIN", "SHOW"]
        every = ok(agents, ex, "timesheets.projects.list", {"include_archived": True})["projects"]
        assert [p["code"] for p in every] == ["ADMIN", "MOVE", "SHOW"]


class TestLogging:
    def test_hours_by_number_or_by_the_clock(self, agents):
        ex, _ = make()
        projects(agents, ex)
        first = log(agents, ex, "SHOW", 7, 2.5, "Layout review")
        assert (first["entry"]["hours"], first["entry"]["week"]) == (2.5, WEEK)
        second = ok(agents, ex, "timesheets.entries.log", {
            "project_code": "show", "date": "2026-09-07", "start": "13:00", "end": "14:40"})
        assert second["entry"]["hours"] == 1.75 and second["day_total"] == 4.25

    def test_a_day_holds_at_most_24_hours(self, agents):
        ex, _ = make()
        projects(agents, ex)
        log(agents, ex, "SHOW", 7, 20)
        assert "25.00" in refused(agents, ex, "timesheets.entries.log", {
            "project_code": "ADMIN", "date": "2026-09-07", "hours": 5})

    def test_an_unknown_project_names_the_known_ones(self, agents):
        ex, _ = make()
        projects(agents, ex)
        error = refused(agents, ex, "timesheets.entries.log", {"project_code": "XYZ", "hours": 1})
        assert "XYZ" in error and "SHOW" in error


class TestFromTheCalendar:
    def test_rules_place_meetings_and_the_rest_come_back(self, agents):
        ex, _ = make()
        projects(agents, ex)
        result = ok(agents, ex, "timesheets.entries.from_events", {"events": CALENDAR})
        assert [(c["event_id"], c["project_code"], c["hours"], c["date"]) for c in result["created"]] == [
            ("ev1", "ADMIN", 0.5, "2026-09-07"),
            ("ev2", "SHOW", 1.75, "2026-09-07"),
            ("ev3", "MOVE", 1.0, "2026-09-08")]
        unassigned = {u["event_id"]: u["reason"] for u in result["unassigned"]}
        assert unassigned["ev4"] == "no rule places it"
        assert "more than one project" in unassigned["ev5"]
        assert [s["event_id"] for s in result["skipped"]] == ["ev6", "ev7"]

    def test_a_meeting_is_never_logged_twice_and_answers_place_the_rest(self, agents):
        ex, _ = make()
        projects(agents, ex)
        ok(agents, ex, "timesheets.entries.from_events", {"events": CALENDAR})
        again = ok(agents, ex, "timesheets.entries.from_events", {
            "events": CALENDAR, "assign": [{"event_id": "ev4", "project_code": "admin"},
                                           {"event_id": "ev5", "project_code": "SHOW"}]})
        assert [c["event_id"] for c in again["created"]] == ["ev4", "ev5"]
        assert again["unassigned"] == []
        already = [s["event_id"] for s in again["skipped"] if s["reason"] == "already on the timesheet"]
        assert already == ["ev1", "ev2", "ev3"]
        listed = ok(agents, ex, "timesheets.entries.list", {"week": WEEK})
        assert listed["total"] == 0.5 + 1.75 + 1.0 + 1.0 + 0.25


class TestTheWeek:
    def full_week(self, agents, ex):
        projects(agents, ex)
        log(agents, ex, "SHOW", 7, 6)
        log(agents, ex, "ADMIN", 7, 2)
        log(agents, ex, "SHOW", 8, 8)
        log(agents, ex, "MOVE", 9, 3)
        log(agents, ex, "SHOW", 10, 8)

    def test_the_summary_adds_up_and_finds_the_gaps_that_have_happened(self, agents):
        ex, _ = make()
        self.full_week(agents, ex)
        week = ok(agents, ex, "timesheets.weeks.summary", {"week": WEEK, "as_of": "2026-09-10"})
        assert (week["from"], week["to"], week["status"]) == ("2026-09-07", "2026-09-13", "open")
        assert [p["project_code"] for p in week["projects"]] == ["ADMIN", "MOVE", "SHOW"]
        show = next(p for p in week["projects"] if p["project_code"] == "SHOW")
        assert show["days"] == [6.0, 8.0, 0.0, 8.0, 0.0, 0.0, 0.0] and show["total"] == 22.0
        assert week["day_totals"] == [8.0, 8.0, 3.0, 8.0, 0.0, 0.0, 0.0]
        assert (week["total"], week["billable_hours"]) == (27.0, 22.0)
        assert week["gaps"] == [{"date": "2026-09-09", "logged": 3.0, "short": 5.0}]
        later = ok(agents, ex, "timesheets.weeks.summary", {"week": WEEK, "as_of": "2026-09-13"})
        assert [g["date"] for g in later["gaps"]] == ["2026-09-09", "2026-09-11"]
        budget = next(p for p in ok(agents, ex, "timesheets.projects.list", {})["projects"]
                      if p["code"] == "SHOW")
        assert (budget["logged_hours"], budget["budget_left"]) == (22.0, 38.0)

    def test_a_submitted_week_is_locked_until_reopened(self, agents):
        ex, provider = make()
        self.full_week(agents, ex)
        entry_ref = ok(agents, ex, "timesheets.entries.list", {"week": WEEK})["entries"][0]["entry_ref"]
        submitted = ok(agents, ex, "timesheets.weeks.submit", {"week": WEEK, "as_of": "2026-09-11"})
        assert (submitted["status"], submitted["total"], submitted["entries"]) == ("submitted", 27.0, 5)
        assert [g["date"] for g in submitted["gaps"]] == ["2026-09-09", "2026-09-11"]
        assert "submitted" in refused(agents, ex, "timesheets.entries.log", {
            "project_code": "SHOW", "date": "2026-09-11", "hours": 4})
        assert "submitted" in refused(agents, ex, "timesheets.entries.update",
                                      {"entry_ref": entry_ref, "hours": 1})
        assert "submitted" in refused(agents, ex, "timesheets.entries.remove", {"entry_ref": entry_ref})
        late = ok(agents, ex, "timesheets.entries.from_events", {
            "events": [meeting("ev9", "Harbourline follow-up", 11, "10:00", "11:00")]})
        assert late["created"] == [] and "submitted" in late["skipped"][0]["reason"]
        assert "already" in refused(agents, ex, "timesheets.weeks.submit", {"week": WEEK})
        ok(agents, ex, "timesheets.weeks.reopen", {"week": WEEK})
        log(agents, ex, "SHOW", 11, 4)
        statuses = {e["status"] for e in ok(agents, ex, "timesheets.entries.list", {"week": WEEK})["entries"]}
        assert statuses == {"draft"}

    def test_the_workbook_reads_back(self, agents):
        ex, provider = make()
        self.full_week(agents, ex)
        made = ok(agents, ex, "timesheets.weeks.workbook", {"week": WEEK}, chat_level=2)
        assert (made["filename"], made["total"], made["verified"]) == ("timesheet-2026-W37.xlsx", 27.0, True)
        raw = provider.files["timesheets__workbook"][made["file_ref"]]["content"]
        wb = load_workbook(io.BytesIO(raw))
        assert wb.sheetnames == ["Timesheet", "Entries"]
        grid = list(wb["Timesheet"].iter_rows(values_only=True))
        assert grid[2][:5] == ("Project", "Name", "Billable", "Mon 07", "Tue 08")
        assert grid[-1][0] == "Total" and grid[-1][-1] == 27
        assert len(list(wb["Entries"].iter_rows(values_only=True))) == 1 + 5

    def test_an_empty_week_has_nothing_to_hand_in(self, agents):
        ex, _ = make()
        projects(agents, ex)
        assert "Nothing" in refused(agents, ex, "timesheets.weeks.submit", {"week": WEEK})
        assert "ISO week" in refused(agents, ex, "timesheets.weeks.summary", {"week": "37"})
