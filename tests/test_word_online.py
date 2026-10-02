"""The Word Online agent in a real worker against a loopback Microsoft
Graph serving real .docx files built here with python-docx: Sidra Office
Supplies' proposal for Harbourline — a heading, a delivery paragraph in
two runs with a comment from Dana on it, a price table — and a returns
policy.
"""

import asyncio
import base64
import io
import random

import docx
import pytest
from docx.oxml.ns import qn

from ai_runtime.execution.executor import FunctionExecutor
from sim.resources import InMemoryResourceProvider

from tests.graph_word_stub import GraphWordStub

DANA = "dana@harbourline.example"
DELIVERY = "Delivery takes fifteen working days from the signed order."
TEN_DAYS = "Delivery takes ten working days from the signed order."
SCOPES = ["openid", "email", "offline_access", "User.Read", "Mail.ReadWrite", "Mail.Send",
          "Calendars.ReadWrite", "Files.ReadWrite.All", "Chat.ReadWrite", "Team.ReadBasic.All",
          "Channel.ReadBasic.All", "ChannelMessage.Send", "Tasks.ReadWrite"]


def run(awaitable):
    return asyncio.run(awaitable)


def proposal(comments=(("Dana Harbour", "Can this be ten days?"),), delivery=DELIVERY,
             extra_paragraph=""):
    document = docx.Document()
    document.add_heading("Harbourline proposal", level=1)
    document.add_paragraph("Sidra Office Supplies will furnish twenty desks.")
    document.add_heading("Delivery", level=2)
    paragraph = document.add_paragraph()
    lead = paragraph.add_run(delivery[:15])
    lead.bold = True
    paragraph.add_run(delivery[15:])
    table = document.add_table(rows=2, cols=2)
    for row, (item, price) in zip(table.rows, (("Desk", "120.00"), ("Chair", "45.50"))):
        row.cells[0].text, row.cells[1].text = item, price
    document.add_paragraph("Prices are valid until 30 September 2026.")
    if extra_paragraph:
        document.add_paragraph(extra_paragraph)
    for number, (author, text) in enumerate(comments):
        comment = document.add_comment(paragraph.runs, text=text, author=author,
                                       initials="".join(w[0] for w in author.split()))
        # Word keeps a comment's date when the file is saved again; the
        # rebuilt fixtures must too, or every save would look like news.
        comment._comment_elm.set(qn("w:date"), f"2026-09-0{number + 1}T10:00:00Z")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def returns_policy():
    document = docx.Document()
    document.add_paragraph("Damaged goods are replaced free of charge.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def paragraphs_of(raw):
    return [p for p in docx.Document(io.BytesIO(raw)).paragraphs]


@pytest.fixture
def graph():
    stub = GraphWordStub().start()
    stub.returns = stub.add_file("Sidra returns policy.docx", returns_policy())
    stub.add_file("Harbourline proposal.pdf", b"%PDF-1.4 exported")
    stub.proposal = stub.add_file("Harbourline proposal.docx", proposal(), by=DANA)
    yield stub
    stub.stop()


def executor(graph, access_token="at-1"):
    provider = InMemoryResourceProvider(
        secrets={"word_online__microsoft": graph.secret(access_token)})
    return FunctionExecutor(provider=provider), provider


def invoke(agents, ex, name, inputs, chat_level=0):
    return run(ex.invoke(agents["word_online"], name, inputs, chat_level=chat_level))


class TestThePackage:
    def test_the_worker_verifies_every_declared_function(self, agents):
        from ai_runtime.agents.worker_handle import WorkerHandle

        agent = agents["word_online"]
        assert WorkerHandle.probe(agent.environment.python, agent.folder,
                                  agent.manifest.document) == []

    def test_it_asks_microsoft_for_the_shared_scope_list(self, agents):
        oauth = agents["word_online"].manifest.resource("secrets", "microsoft")["oauth"]
        assert oauth["scopes"] == SCOPES

    def test_what_each_function_costs(self, agents):
        levels = {f"{t['id']}.{f['id']}": f["permission_level"]
                  for t in agents["word_online"].manifest.document["tools"]
                  for f in t["functions"]}
        assert levels == {
            "account.status": 0, "docs.find": 0, "docs.read": 0, "comments.list": 0,
            "edits.propose": 1, "edits.apply": 3, "edits.discard": 1,
            "changes.watch": 1, "changes.new": 1}
        _, new = agents["word_online"].manifest.function("word_online.changes.new")
        assert new["schedulable"] is True and not new.get("llm")

    def test_status_reports_the_account(self, agents, graph):
        ex, _ = executor(graph)
        result, status = invoke(agents, ex, "word_online.account.status", {})
        assert status == "success", result
        assert result == {"connected": True, "email": GraphWordStub.ACCOUNT}


class TestReading:
    def test_find_returns_only_word_files_newest_first(self, agents, graph):
        graph.add_file("Harbourline proposal (old).docx", returns_policy())
        ex, _ = executor(graph)
        result, status = invoke(agents, ex, "word_online.docs.find", {"query": "harbourline"})
        assert status == "success", result
        assert [d["name"] for d in result["documents"]] == [
            "Harbourline proposal (old).docx", "Harbourline proposal.docx"]

    def test_a_document_is_read_as_numbered_paragraphs_with_its_etag(self, agents, graph):
        ex, _ = executor(graph)
        first, status = invoke(agents, ex, "word_online.docs.read",
                               {"item_id": graph.proposal, "max_paragraphs": 4})
        assert status == "success", first
        assert first["etag"] == graph.etag(graph.proposal)
        assert (first["modified_by"], first["total_paragraphs"], first["comment_count"]) == (DANA, 9, 1)
        assert [(p["index"], p["style"], p["text"]) for p in first["paragraphs"]] == [
            (1, "Heading 1", "Harbourline proposal"),
            (2, "Normal", "Sidra Office Supplies will furnish twenty desks."),
            (3, "Heading 2", "Delivery"), (4, "Normal", DELIVERY)]
        assert first["next_from"] == 5
        rest, _ = invoke(agents, ex, "word_online.docs.read",
                         {"item_id": graph.proposal, "from": 5})
        assert [(p["text"], p["in_table"]) for p in rest["paragraphs"]] == [
            ("Desk", True), ("120.00", True), ("Chair", True), ("45.50", True),
            ("Prices are valid until 30 September 2026.", False)]
        assert "next_from" not in rest

    def test_comments_are_read_from_the_file_with_the_paragraph_they_sit_on(self, agents, graph):
        ex, _ = executor(graph)
        result, status = invoke(agents, ex, "word_online.comments.list", {"item_id": graph.proposal})
        assert status == "success", result
        [comment] = result["comments"]
        assert (comment["author"], comment["text"], comment["paragraph"], comment["paragraph_text"]) == (
            "Dana Harbour", "Can this be ten days?", 4, DELIVERY)
        assert result["total"] == 1 and "next_from" not in result

    def test_a_file_that_is_not_word_is_named_for_what_it_is(self, agents, graph):
        pdf = next(i for i, item in graph.items.items() if item["name"].endswith(".pdf"))
        ex, _ = executor(graph)
        result, status = invoke(agents, ex, "word_online.docs.read", {"item_id": pdf})
        assert status == "error" and "not a Word" in result["error"]


class TestEdits:
    def propose(self, agents, ex, graph, paragraph=4, text=TEN_DAYS):
        result, status = invoke(agents, ex, "word_online.edits.propose", {
            "item_id": graph.proposal, "paragraph": paragraph, "new_text": text}, chat_level=1)
        assert status == "success", result
        return result

    def test_a_proposal_is_recorded_and_the_file_does_not_move(self, agents, graph):
        ex, provider = executor(graph)
        before = graph.etag(graph.proposal)
        proposed = self.propose(agents, ex, graph)
        assert (proposed["old_text"], proposed["etag"]) == (DELIVERY, before)
        keys = provider.data["word_online__edit"][proposed["edit_ref"]]["keys"]
        assert (keys["paragraph"], keys["status"], keys["etag"]) == (4, "proposed", before)
        assert graph.puts == [] and graph.etag(graph.proposal) == before

    def test_applying_replaces_the_words_and_keeps_the_formatting_and_the_comment(self, agents, graph):
        ex, provider = executor(graph)
        proposed = self.propose(agents, ex, graph)
        applied, status = invoke(agents, ex, "word_online.edits.apply",
                                 {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "success", applied
        [put] = graph.puts
        assert put["if_match"] == proposed["etag"]
        assert applied["etag"] == graph.etag(graph.proposal) != proposed["etag"]
        saved = paragraphs_of(graph.items[graph.proposal]["content"])
        assert saved[3].text == TEN_DAYS and saved[3].runs[0].bold is True
        assert [p.text for p in saved[:3]] == [
            "Harbourline proposal", "Sidra Office Supplies will furnish twenty desks.", "Delivery"]
        assert provider.data["word_online__edit"][proposed["edit_ref"]]["keys"]["status"] == "applied"

        comments, _ = invoke(agents, ex, "word_online.comments.list", {"item_id": graph.proposal})
        assert [(c["text"], c["paragraph"], c["paragraph_text"]) for c in comments["comments"]] == [
            ("Can this be ten days?", 4, TEN_DAYS)]

        again, status = invoke(agents, ex, "word_online.edits.apply",
                               {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "error" and "applied" in again["error"]
        assert len(graph.puts) == 1

    def test_an_edit_to_a_file_saved_since_is_refused_not_clobbered(self, agents, graph):
        ex, provider = executor(graph)
        proposed = self.propose(agents, ex, graph)
        dana_version = proposal(extra_paragraph="Installation is included.")
        graph.save_as(graph.proposal, dana_version, DANA)
        downloads = len(graph.downloads)
        result, status = invoke(agents, ex, "word_online.edits.apply",
                                {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "error" and result["kind"] == "conflict", result
        assert DANA in result["error"] and "changed" in result["error"]
        assert graph.puts == [] and len(graph.downloads) == downloads
        assert graph.items[graph.proposal]["content"] == dana_version
        keys = provider.data["word_online__edit"][proposed["edit_ref"]]["keys"]
        assert keys["status"] == "refused" and "changed" in keys["reason"]

    def test_a_save_landing_between_download_and_upload_is_refused_by_onedrive(self, agents, graph):
        ex, provider = executor(graph)
        proposed = self.propose(agents, ex, graph)
        dana_version = proposal(delivery="Delivery is by arrangement.")
        graph.before_put = lambda item_id: graph.save_as(item_id, dana_version, DANA)
        result, status = invoke(agents, ex, "word_online.edits.apply",
                                {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "error" and result["kind"] == "conflict", result
        assert graph.items[graph.proposal]["content"] == dana_version
        assert provider.data["word_online__edit"][proposed["edit_ref"]]["keys"]["status"] == "refused"

    def test_a_save_with_no_answer_is_unknown_and_not_sent_again(self, agents, graph):
        ex, provider = executor(graph)
        proposed = self.propose(agents, ex, graph)
        graph.drop_put = True
        result, status = invoke(agents, ex, "word_online.edits.apply",
                                {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "error" and result["kind"] == "unknown"
        again, status = invoke(agents, ex, "word_online.edits.apply",
                               {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "error" and again["kind"] == "unknown"
        assert len(graph.puts) == 1
        assert provider.data["word_online__edit"][proposed["edit_ref"]]["keys"]["status"] == "unknown"

    def test_a_document_too_large_to_carry_is_refused_with_its_size(self, agents, graph):
        """Refused where it is first read, before anything is downloaded
        or proposed."""
        noise = base64.b64encode(random.Random(7).randbytes(30_000_000)).decode("ascii")
        big = graph.add_file("Catalogue.docx", proposal(extra_paragraph=noise))
        assert len(graph.items[big]["content"]) > 25 * 1024 * 1024
        ex, provider = executor(graph)
        downloads = len(graph.downloads)
        result, status = invoke(agents, ex, "word_online.edits.propose", {
            "item_id": big, "paragraph": 4, "new_text": TEN_DAYS}, chat_level=1)
        assert status == "error" and result["kind"] == "too_large", result
        assert f"{len(graph.items[big]['content']) // 1024:,} KB" in result["error"]
        assert graph.puts == [] and len(graph.downloads) == downloads
        assert not provider.data.get("word_online__edit")

    def test_applying_needs_the_chats_trust_and_discard_leaves_the_file(self, agents, graph):
        ex, _ = executor(graph)
        proposed = self.propose(agents, ex, graph)
        denied, status = invoke(agents, ex, "word_online.edits.apply",
                                {"edit_ref": proposed["edit_ref"]}, chat_level=1)
        assert status == "error" and "approv" in denied["error"].lower()
        discarded, status = invoke(agents, ex, "word_online.edits.discard",
                                   {"edit_ref": proposed["edit_ref"]}, chat_level=1)
        assert status == "success" and discarded["discarded"] is True
        result, status = invoke(agents, ex, "word_online.edits.apply",
                                {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "error" and "discarded" in result["error"]
        assert graph.puts == []

    def test_a_paragraph_number_past_the_end_records_nothing(self, agents, graph):
        ex, provider = executor(graph)
        result, status = invoke(agents, ex, "word_online.edits.propose", {
            "item_id": graph.proposal, "paragraph": 10, "new_text": "x"}, chat_level=1)
        assert status == "error" and "9 paragraphs" in result["error"]
        assert provider.data == {}


class TestWatching:
    def test_the_watch_reports_a_change_and_its_new_comments_exactly_once(self, agents, graph):
        ex, _ = executor(graph)
        watch, status = invoke(agents, ex, "word_online.changes.watch",
                               {"item_id": graph.proposal}, chat_level=1)
        assert status == "success", watch
        assert watch["comment_count"] == 1

        downloads = len(graph.downloads)
        quiet, status = invoke(agents, ex, "word_online.changes.new", {}, chat_level=1)
        assert status == "success", quiet
        assert (quiet["checked"], quiet["changes"], quiet["comments"]) == (1, [], [])
        assert len(graph.downloads) == downloads      # an unchanged eTag downloads nothing

        graph.save_as(graph.proposal, proposal(comments=(
            ("Dana Harbour", "Can this be ten days?"),
            ("Dana Harbour", "Also: can you deliver chairs first?"))), DANA)
        news, _ = invoke(agents, ex, "word_online.changes.new", {}, chat_level=1)
        [change] = news["changes"]
        assert (change["watch_ref"], change["modified_by"], change["new_comments"]) == (
            watch["watch_ref"], DANA, 1)
        [comment] = news["comments"]
        assert (comment["text"], comment["paragraph"]) == ("Also: can you deliver chairs first?", 4)

        again, _ = invoke(agents, ex, "word_online.changes.new", {}, chat_level=1)
        assert again["changes"] == [] and again["comments"] == []

    def test_the_accounts_own_saves_and_comments_are_not_news(self, agents, graph):
        ex, _ = executor(graph)
        watch, _ = invoke(agents, ex, "word_online.changes.watch",
                          {"item_id": graph.proposal}, chat_level=1)
        proposed, _ = invoke(agents, ex, "word_online.edits.propose", {
            "item_id": graph.proposal, "paragraph": 4, "new_text": TEN_DAYS}, chat_level=1)
        applied, status = invoke(agents, ex, "word_online.edits.apply",
                                 {"edit_ref": proposed["edit_ref"]}, chat_level=3)
        assert status == "success", applied
        quiet, _ = invoke(agents, ex, "word_online.changes.new",
                          {"watch_ref": watch["watch_ref"]}, chat_level=1)
        assert quiet["changes"] == [] and quiet["comments"] == []

        graph.save_as(graph.proposal, proposal(comments=(
            ("Dana Harbour", "Can this be ten days?"),
            (GraphWordStub.DISPLAY_NAME, "Checking with the warehouse."))), GraphWordStub.ACCOUNT)
        quiet, _ = invoke(agents, ex, "word_online.changes.new", {}, chat_level=1)
        assert quiet["changes"] == [] and quiet["comments"] == []

        # Someone else's save with no comment is still a change.
        graph.save_as(graph.proposal, proposal(extra_paragraph="Installation is included."), DANA)
        news, _ = invoke(agents, ex, "word_online.changes.new", {}, chat_level=1)
        assert [(c["modified_by"], c["new_comments"]) for c in news["changes"]] == [(DANA, 0)]


class TestTheConnection:
    def test_an_expired_token_says_reconnect_and_writes_nothing(self, agents, graph):
        ex, provider = executor(graph, access_token="expired")
        result, status = invoke(agents, ex, "word_online.docs.read", {"item_id": graph.proposal})
        assert status == "error" and result["kind"] == "auth"
        assert "reconnect" in result["error"]
        for name, inputs in (("word_online.edits.propose",
                              {"item_id": graph.proposal, "paragraph": 4, "new_text": "x"}),
                             ("word_online.changes.watch", {"item_id": graph.proposal})):
            result, status = invoke(agents, ex, name, inputs, chat_level=1)
            assert status == "error" and result["kind"] == "auth", name
        assert provider.data == {} and graph.puts == []

    def test_no_credential_is_a_plain_answer(self, agents):
        ex = FunctionExecutor(provider=InMemoryResourceProvider())
        result, status = invoke(agents, ex, "word_online.account.status", {})
        assert status == "success" and result["connected"] is False
        assert "No Microsoft account" in result["problem"]
