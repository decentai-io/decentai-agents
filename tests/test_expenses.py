"""The Expenses agent in a real worker: Omar's trip receipts for
Sidra Office Supplies — text receipts, one photo, one duplicate, one over the
meal allowance, one in another currency — and the company policy."""

import asyncio
import io
import json

from openpyxl import load_workbook

from ai_runtime.execution.executor import FunctionExecutor
from sim.resources import InMemoryResourceProvider

RECEIPTS = {
    "lunch-0907.txt": "SHORELINE CAFE\nDubai Marina\n07 Sep 2026 13:12\nGrilled fish 48.00\nWater 7.00\nTOTAL AED 55.00\nPaid card",
    "dinner-0907.txt": "The Copper Pot\n2026-09-07\n2 x main course 140.00\nDessert 32.00\nService 17.20\nTotal USD 189.20",
    "taxi-0908.txt": "CAREEM RECEIPT\n8 September 2026\nTrip: Marina -> Harbourline offices\nFare USD 24.50\nTotal USD 24.50",
    "hotel-0908.txt": "Harbour View Hotel\nCheck-in 07/09/2026  Check-out 08/09/2026\nRoom 1 night 210.00\nCity tax 12.00\nAmount due USD 222.00",
    "dinner-0907-again.txt": "The Copper Pot\n2026-09-07\n2 x main course 140.00\nDessert 32.00\nService 17.20\nTotal USD 189.20",
    "stationery-0909.txt": "OfficeMart\nSep 9, 2026\nNotebooks 18.00\nPens 6.50\nTotal 24.50",
}
PHOTO = b"\xff\xd8\xff\xe0\x00\x10JFIF fictional photo bytes"
POLICY = """Sidra Office Supplies Expense Policy (fictional)

1. Meals: employees may claim up to USD 60 per person per day for meals while travelling.
2. Hotels: accommodation up to USD 200 per night; anything above needs prior approval.
3. Receipts are required for any expense above USD 25.
4. Claims must be submitted within 30 days of the expense date.
5. Alcohol is not reimbursable.
"""

MODEL_ANSWERS = {
    "lunch-0907.txt": {"merchant": "Shoreline Cafe", "date": "07 Sep 2026", "currency": "AED",
                       "amount": "55.00", "category": "meals", "quote_amount": "TOTAL AED 55.00"},
    "dinner-0907.txt": {"merchant": "The Copper Pot", "date": "2026-09-07", "currency": "USD",
                        "amount": "189.20", "category": "meals", "quote_amount": "Total USD 189.20"},
    "taxi-0908.txt": {"merchant": "Careem", "date": "8 September 2026", "currency": "USD",
                      "amount": "24.50", "category": "transport", "quote_amount": "Total USD 24.50"},
    "hotel-0908.txt": {"merchant": "Harbour View Hotel", "date": "08/09/2026", "currency": "USD",
                       "amount": "222.00", "category": "accommodation", "quote_amount": "Amount due USD 222.00"},
    "dinner-0907-again.txt": {"merchant": "The Copper Pot", "date": "2026-09-07", "currency": "USD",
                              "amount": "189.20", "category": "meals", "quote_amount": "Total USD 189.20"},
    # The model misreads the stationery total and names no currency.
    "stationery-0909.txt": {"merchant": "OfficeMart", "date": "Sep 9, 2026", "currency": "",
                            "amount": "25.40", "category": "office", "quote_amount": "Total 24.50"},
}
POLICY_ANSWER = {
    "meal_max_per_day": {"value": "60", "quote": "employees may claim up to USD 60 per person per day for meals while travelling."},
    "hotel_max_per_night": {"value": "200", "quote": "accommodation up to USD 200 per night"},
    "receipt_required_above": {"value": "25", "quote": "Receipts are required for any expense above USD 25."},
    "claim_within_days": {"value": "30", "quote": "Claims must be submitted within 30 days of the expense date."},
    "alcohol_allowed": {"value": "no", "quote": "Alcohol is never ever reimbursable."},   # misquoted
    "taxi_max_per_trip": {"found": False},
}


async def model(messages, max_tokens=None):
    prompt = messages[-1]["content"]
    if prompt.startswith("POLICY"):
        return json.dumps(POLICY_ANSWER)
    for name, answer in MODEL_ANSWERS.items():
        if f"RECEIPT ({name})" in prompt:
            return json.dumps(answer)
    return "{}"


def run(awaitable):
    return asyncio.run(awaitable)


def make():
    provider = InMemoryResourceProvider()
    return FunctionExecutor(provider=provider, llm=model), provider


def invoke(agents, ex, name, inputs, chat_level=1):
    return run(ex.invoke(agents["expenses"], name, inputs, chat_level=chat_level))


def upload(provider, filename, raw: bytes) -> str:
    return run(provider.create_file("chat_attachment", filename, raw))["resource_ref"]


def trip(agents, ex, provider, names=None):
    claim, _ = invoke(agents, ex, "expenses.claims.create", {"title": "Dubai trip, 7–9 Sep", "claimant": "Omar"})
    read = {}
    for name in (RECEIPTS if names is None else names):
        ref = upload(provider, name, RECEIPTS[name].encode())
        read[name], status = invoke(agents, ex, "expenses.receipts.read",
                                    {"claim_ref": claim["claim_ref"], "file_ref": ref})
        assert status == "success", read[name]
    return claim["claim_ref"], read


class TestThePackage:
    def test_the_worker_verifies_every_declared_function(self, agents):
        from ai_runtime.agents.worker_handle import WorkerHandle

        agent = agents["expenses"]
        assert WorkerHandle.probe(agent.environment.python, agent.folder,
                                  agent.manifest.document) == []


class TestReceipts:
    def test_facts_are_verified_against_the_receipt(self, agents):
        ex, provider = make()
        _, read = trip(agents, ex, provider, ["dinner-0907.txt", "taxi-0908.txt", "hotel-0908.txt"])
        dinner = read["dinner-0907.txt"]
        assert dinner["status"] == "extracted"
        assert dinner["amount"] == 189.2 and dinner["currency"] == "USD" and dinner["date"] == "2026-09-07"
        assert set(dinner["verified"]) == {"merchant", "date", "amount", "currency"}
        assert dinner["assumed"] == ["category"]
        taxi = read["taxi-0908.txt"]
        assert taxi["date"] == "2026-09-08" and taxi["category"] == "transport"
        assert "merchant" in taxi["verified"]             # "Careem" appears in CAREEM RECEIPT
        # The hotel's dates are day/month-ambiguous (07/09/2026): not read,
        # never guessed — the receipt needs review for its date.
        hotel = read["hotel-0908.txt"]
        assert hotel["status"] == "needs_review" and hotel["date"] == ""
        assert "date not readable" in hotel["note"] and hotel["amount"] == 222.0

    def test_a_misread_amount_and_a_missing_currency_need_review(self, agents):
        ex, provider = make()
        _, read = trip(agents, ex, provider, ["stationery-0909.txt"])
        item = read["stationery-0909.txt"]
        assert item["status"] == "needs_review"
        assert "amount" in item["assumed"]                # 25.40 is not on the receipt
        assert item["currency"] == "" and "currency not shown" in item["note"]
        assert item["date"] == "2026-09-09"

    def test_a_photo_is_manual_and_confirm_fills_it(self, agents):
        ex, provider = make()
        claim_ref, _ = trip(agents, ex, provider, [])
        ref = upload(provider, "parking.jpg", PHOTO)
        result, status = invoke(agents, ex, "expenses.receipts.read", {"claim_ref": claim_ref, "file_ref": ref})
        assert status == "success" and result["status"] == "manual"
        assert "photo" in result["note"]
        refused, status = invoke(agents, ex, "expenses.receipts.confirm", {
            "receipt_ref": result["receipt_ref"], "merchant": "Marina Parking", "amount": 8})
        assert status == "error" and "date" in refused["error"]
        confirmed, status = invoke(agents, ex, "expenses.receipts.confirm", {
            "receipt_ref": result["receipt_ref"], "merchant": "Marina Parking", "date": "2026-09-08",
            "currency": "usd", "amount": 8, "category": "transport"})
        assert status == "success", confirmed
        assert confirmed["status"] == "confirmed" and confirmed["currency"] == "USD"
        got, _ = invoke(agents, ex, "expenses.claims.get", {"claim_ref": claim_ref})
        assert got["totals"] == [{"currency": "USD", "amount": 8.0, "receipts": 1}]
        assert got["needs_attention"] == 0

    def test_a_duplicate_is_seen_on_read(self, agents):
        ex, provider = make()
        _, read = trip(agents, ex, provider, ["dinner-0907.txt", "dinner-0907-again.txt"])
        again = read["dinner-0907-again.txt"]
        assert again["status"] == "duplicate"
        assert again["duplicate_of"] == read["dinner-0907.txt"]["receipt_ref"]


class TestPolicy:
    def test_rules_come_with_verified_quotes_and_honest_gaps(self, agents):
        ex, provider = make()
        ref = upload(provider, "expense-policy.txt", POLICY.encode())
        result, status = invoke(agents, ex, "expenses.policies.load", {"file_ref": ref})
        assert status == "success", result
        by_name = {r["name"]: r for r in result["rules"]}
        assert by_name["meal_max_per_day"]["value"] == "60" and by_name["meal_max_per_day"]["verified"] is True
        assert by_name["alcohol_allowed"]["verified"] is False       # the quote is not in the policy
        assert result["not_found"] == ["taxi_max_per_trip"]
        assert result["unverified"] == ["alcohol_allowed"]
        # A person sets what the policy lacks, and it replaces nothing else.
        set_, _ = invoke(agents, ex, "expenses.policies.set_rule", {"name": "taxi_max_per_trip", "value": "40"})
        rules, _ = invoke(agents, ex, "expenses.policies.get", {})
        assert {r["name"]: r["value"] for r in rules["rules"]}["taxi_max_per_trip"] == "40"
        assert len(rules["rules"]) == 6
        # Loading again replaces, never duplicates.
        invoke(agents, ex, "expenses.policies.load", {"file_ref": ref})
        rules, _ = invoke(agents, ex, "expenses.policies.get", {})
        assert len(rules["rules"]) == 6


class TestTheClaim:
    def test_a_claim_that_exists_is_found_by_listing(self, agents):
        ex, provider = make()
        claim_ref, _ = trip(agents, ex, provider, names=[])
        result, status = invoke(agents, ex, "expenses.claims.list", {})
        assert status == "success", result
        assert [(c["claim_ref"], c["title"], c["status"]) for c in result["claims"]] == [
            (claim_ref, "Dubai trip, 7–9 Sep", "draft")]
        none, _ = invoke(agents, ex, "expenses.claims.list", {"status": "paid"})
        assert none["claims"] == []

    def test_check_flags_duplicates_allowance_currency_and_review(self, agents):
        ex, provider = make()
        policy = upload(provider, "expense-policy.txt", POLICY.encode())
        invoke(agents, ex, "expenses.policies.load", {"file_ref": policy})
        claim_ref, read = trip(agents, ex, provider)
        result, status = invoke(agents, ex, "expenses.claims.check", {"claim_ref": claim_ref, "as_of": "2026-09-12"})
        assert status == "success", result
        kinds = {(e["filename"], e["kind"]) for e in result["exceptions"]}
        assert ("dinner-0907-again.txt", "duplicate") in kinds
        assert ("lunch-0907.txt", "currency") in kinds                  # AED, not converted
        assert ("dinner-0907.txt", "policy") in kinds                    # 189.20 > 60 meals per day
        assert ("hotel-0908.txt", "policy") in kinds                     # 222 > 200 per night
        assert ("stationery-0909.txt", "needs_review") in kinds
        assert result["receipts"] == 6 and result["ok"] == 1            # only the taxi
        assert "meal_max_per_day = 60" in result["rules_applied"]
        detail = next(e["detail"] for e in result["exceptions"]
                      if e["filename"] == "hotel-0908.txt" and e["kind"] == "policy")
        assert detail == "accommodation 222.0 exceeds hotel_max_per_night 200.0"

    def test_the_workbook_and_the_status_rules(self, agents):
        ex, provider = make()
        claim_ref, read = trip(agents, ex, provider, ["dinner-0907.txt", "taxi-0908.txt", "stationery-0909.txt"])
        made, status = invoke(agents, ex, "expenses.claims.workbook",
                              {"claim_ref": claim_ref, "as_of": "2026-09-12"}, chat_level=2)
        assert status == "success", made
        assert made["receipts"] == 3 and made["exceptions"] == 1 and made["verified"] is True
        raw = provider.files["expenses__workbook"][made["file_ref"]]["content"]
        wb = load_workbook(io.BytesIO(raw))
        assert wb.sheetnames == ["Receipts", "Summary", "Exceptions", "Rules applied"]
        rows = list(wb["Receipts"].iter_rows(values_only=True))
        assert rows[3][:6] == ("Date", "Merchant", "Category", "Currency", "Amount", "Status")
        assert any(r[1] == "OfficeMart" and r[5] == "needs_review" for r in rows[4:])
        summary = list(wb["Summary"].iter_rows(values_only=True))
        assert ("meals", "USD", 1, 189.2) in summary

        # Cannot submit while a receipt needs review; can after it is confirmed.
        refused, status = invoke(agents, ex, "expenses.claims.set_status", {"claim_ref": claim_ref, "status": "submitted"})
        assert status == "error" and "stationery-0909.txt" in refused["error"]
        invoke(agents, ex, "expenses.receipts.confirm", {
            "receipt_ref": read["stationery-0909.txt"]["receipt_ref"], "currency": "USD", "amount": 24.5})
        submitted, status = invoke(agents, ex, "expenses.claims.set_status", {"claim_ref": claim_ref, "status": "submitted"})
        assert status == "success" and submitted["status"] == "submitted"
        # Approval needs a reviewer's name; paid needs approval first.
        refused, status = invoke(agents, ex, "expenses.claims.set_status", {"claim_ref": claim_ref, "status": "approved"})
        assert status == "error" and "reviewer" in refused["error"]
        refused, status = invoke(agents, ex, "expenses.claims.set_status", {"claim_ref": claim_ref, "status": "paid", "by": "Finance"})
        assert status == "error" and "cannot become paid" in refused["error"]
        approved, _ = invoke(agents, ex, "expenses.claims.set_status", {
            "claim_ref": claim_ref, "status": "approved", "by": "Dana Whitlock", "on_date": "2026-09-14"})
        assert approved["decided_by"] == "Dana Whitlock" and approved["decided_on"] == "2026-09-14"
        # No more receipts into a submitted claim.
        ref = upload(provider, "late.txt", RECEIPTS["taxi-0908.txt"].encode())
        late, status = invoke(agents, ex, "expenses.receipts.read", {"claim_ref": claim_ref, "file_ref": ref})
        assert status == "error" and "approved" in late["error"]
