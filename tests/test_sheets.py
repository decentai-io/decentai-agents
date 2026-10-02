"""The Spreadsheets agent in a real worker, over two fictional customer
lists from Sidra Office Supplies — one from the accounting export, one from
the sales sheet — with the usual mess: duplicates, a missing phone,
an identifier with leading zeros, a formula, a date."""

import asyncio
import io
from datetime import date

from openpyxl import Workbook, load_workbook

from ai_runtime.execution.executor import FunctionExecutor
from sim.resources import InMemoryResourceProvider

CUSTOMERS_CSV = """Customer ID,Name,Email,Phone,Country
00017,Harbourline Consulting,dana@harbourline.example,+971 50 123 4567,UAE
00018,Cedar Office,buy@cedaroffice.example,,uae
00019,Pier 9 Studio,hello@pier9.example,050-987-6543,UAE
00017,Harbourline Consulting,DANA@harbourline.example,+971501234567,UAE
00020,Northlight Seating,quotes@northlight-seating.example,+971 4 555 0100,U.A.E.
"""


def sales_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Account", "Company", "Contact Email", "Mobile", "Last Order", "Orders", "Value"])
    ws.append(["00017", "Harbourline Consulting", "dana@harbourline.example", "0501234567",
               date(2026, 8, 30), 3, 4120.5])
    ws.append(["00019", "Pier 9 Studio", "hello@pier9.example", "0509876543",
               date(2026, 7, 12), 1, 860])
    ws.append(["00021", "Marina Dental", "office@marinadental.example", "0501112222",
               date(2026, 9, 1), 2, 1299.99])
    ws.append(["00020", "Northlight Seating", "sales@northlight-seating.example", "045550100",
               date(2026, 6, 2), 5, 9800])
    ws.append([None, "Total", None, None, None, "=SUM(F2:F5)", "=SUM(G2:G5)"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def run(awaitable):
    return asyncio.run(awaitable)


def make():
    provider = InMemoryResourceProvider()
    return FunctionExecutor(provider=provider), provider


def invoke(agents, ex, name, inputs, chat_level=2):
    return run(ex.invoke(agents["sheets"], name, inputs, chat_level=chat_level))


def upload(provider, filename, raw: bytes) -> str:
    return run(provider.create_file("chat_attachment", filename, raw))["resource_ref"]


class TestThePackage:
    def test_the_worker_verifies_every_declared_function(self, agents):
        from ai_runtime.agents.worker_handle import WorkerHandle

        agent = agents["sheets"]
        assert WorkerHandle.probe(agent.environment.python, agent.folder,
                                  agent.manifest.document) == []


class TestAskingForColumnsByName:
    """The header a person reads is not a key they typed. A model that
    sees "Customer ID" and asks for "customer id" — or trims, or Title
    Cases — must still get the values."""

    def test_columns_asked_for_in_another_case_still_carry_their_values(self, agents):
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode("utf-8"))
        result, status = invoke(agents, ex, "sheets.read.rows",
                                {"file_ref": ref, "columns": ["CUSTOMER ID", "name"]})
        assert status == "success", result
        assert result["total_rows"] == 5
        # The columns come back under the header's own spelling...
        assert result["columns"] == ["Customer ID", "Name"]
        # ...and every row actually holds them. This used to be five
        # empty dicts: the right number of rows with nothing in them.
        assert len(result["rows"]) == 5
        assert all(row for row in result["rows"]), result["rows"]
        assert result["rows"][0]["Customer ID"] == "00017"
        assert result["rows"][0]["Name"] == "Harbourline Consulting"

    def test_a_name_with_stray_spacing_is_still_the_same_column(self, agents):
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode("utf-8"))
        result, status = invoke(agents, ex, "sheets.read.rows",
                                {"file_ref": ref, "columns": [" Name "]})
        assert status == "success", result
        assert result["columns"] == ["Name"]
        assert result["rows"][0] == {"Name": "Harbourline Consulting"}

    def test_a_column_that_is_not_there_is_still_refused(self, agents):
        """Forgiving about spelling must not mean forgiving about truth."""
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode("utf-8"))
        result, status = invoke(agents, ex, "sheets.read.rows",
                                {"file_ref": ref, "columns": ["Assignee"]})
        assert status == "error"
        assert "Assignee" in result["error"]

    def test_an_empty_page_says_the_sheet_is_not_empty(self, agents):
        """Zero rows from a sheet that has rows is a question about
        from_row, and used to be indistinguishable from an empty file."""
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode("utf-8"))
        result, status = invoke(agents, ex, "sheets.read.rows",
                                {"file_ref": ref, "from_row": 99})
        assert status == "success", result
        assert result["rows"] == []
        assert result["total_rows"] == 5
        assert "has 5" in result["note"]


class TestReading:
    def test_inspect_sees_types_zeros_blanks_and_formulas(self, agents):
        ex, provider = make()
        ref = upload(provider, "sales.xlsx", sales_xlsx())
        result, status = invoke(agents, ex, "sheets.read.inspect", {"file_ref": ref})
        assert status == "success", result
        assert result["kind"] == "xlsx"
        assert result["sheets"] == [{"name": "Sales", "rows": 5}]
        by_name = {c["name"]: c for c in result["columns"] if c["sheet"] == "Sales"}
        assert by_name["Account"]["leading_zeros"] is True and by_name["Account"]["type"] == "text"
        assert by_name["Last Order"]["type"] == "date"
        assert by_name["Value"]["type"] == "number" and by_name["Value"]["formulas"] == 1
        assert by_name["Account"]["blanks"] == 1

        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        result, _ = invoke(agents, ex, "sheets.read.inspect", {"file_ref": ref})
        assert result["kind"] == "csv"
        by_name = {c["name"]: c for c in result["columns"]}
        assert by_name["Customer ID"]["leading_zeros"] is True
        assert by_name["Phone"]["blanks"] == 1

    def test_rows_keep_identifiers_dates_and_formulas(self, agents):
        ex, provider = make()
        ref = upload(provider, "sales.xlsx", sales_xlsx())
        result, status = invoke(agents, ex, "sheets.read.rows", {
            "file_ref": ref, "max_rows": 2, "columns": ["Account", "Last Order", "Value"]})
        assert status == "success", result
        assert result["rows"][0] == {"Account": "00017", "Last Order": "2026-08-30", "Value": "4120.5"}
        assert result["truncated"] is True and result["next_row"] == 3
        rest, _ = invoke(agents, ex, "sheets.read.rows", {
            "file_ref": ref, "from_row": 5, "formulas": True, "columns": ["Company", "Value"]})
        assert rest["rows"] == [{"Company": "Total", "Value": "=SUM(G2:G5)"}]
        bad, status = invoke(agents, ex, "sheets.read.rows", {"file_ref": ref, "columns": ["Nope"]})
        assert status == "error" and "No such column" in bad["error"]


class TestChecking:
    def test_quality_finds_duplicates_gaps_and_variants(self, agents):
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        result, status = invoke(agents, ex, "sheets.check.quality", {
            "file_ref": ref, "key_columns": ["Customer ID"],
            "required_columns": ["Phone", "Email"], "categorical_columns": ["Country"]})
        assert status == "success", result
        assert result["duplicate_groups"] == [{"key": "00017", "rows": [2, 5]}]
        assert result["missing"] == [{"column": "Phone", "rows": [3]}]
        assert result["inconsistent"] == [{"column": "Country", "variants": ["UAE", "uae"]}]
        # A phone-like key compares by digits: two spellings of one number.
        result, _ = invoke(agents, ex, "sheets.check.quality", {
            "file_ref": ref, "key_columns": ["Phone"]})
        assert result["duplicate_groups"] == [{"key": "971501234567", "rows": [2, 5]}]
        job = provider.data["sheets__job"][result["job_ref"]]
        assert job["keys"]["kind"] == "quality"


class TestCleaning:
    def test_dedupe_keeps_the_first_and_logs_the_rest(self, agents):
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        result, status = invoke(agents, ex, "sheets.clean.dedupe", {
            "file_ref": ref, "key_columns": ["Customer ID", "Name"], "filename": "customers-clean"})
        assert status == "success", result
        assert result["kept"] == 4 and result["removed"] == 1 and result["removed_rows"] == [5]
        assert result["filename"] == "customers-clean.xlsx" and result["verified"] is True
        raw = provider.files["sheets__workbook"][result["file_ref"]]["content"]
        wb = load_workbook(io.BytesIO(raw))
        assert wb.sheetnames == ["Clean", "Removed", "Change Log"]
        clean = list(wb["Clean"].iter_rows(values_only=True))
        assert clean[0][:2] == ("Customer ID", "Name")
        assert [r[0] for r in clean[1:]] == ["00017", "00018", "00019", "00020"]   # zeros kept
        removed = list(wb["Removed"].iter_rows(values_only=True))
        assert removed[1][0] == "00017" and removed[1][-1] == 5
        log = list(wb["Change Log"].iter_rows(values_only=True))
        assert log[1][1] == "rules" and log[2] == (5, "removed",
            "duplicate of row 2 on Customer ID, Name: 00017 | harbourline consulting")
        # The original is untouched.
        assert provider.files["chat_attachment"][ref]["content"] == CUSTOMERS_CSV.encode()

    def test_dedupe_preserves_dates_and_formulas(self, agents):
        ex, provider = make()
        ref = upload(provider, "sales.xlsx", sales_xlsx())
        result, _ = invoke(agents, ex, "sheets.clean.dedupe", {
            "file_ref": ref, "key_columns": ["Account"], "filename": "sales-clean"})
        raw = provider.files["sheets__workbook"][result["file_ref"]]["content"]
        ws = load_workbook(io.BytesIO(raw))["Clean"]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[1][4] == date(2026, 8, 30) or str(rows[1][4]).startswith("2026-08-30")
        assert rows[5][6] == "=SUM(G2:G5)"
        assert rows[1][0] == "00017"

    def test_producing_a_workbook_needs_level_two(self, agents):
        ex, provider = make()
        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        result, status = invoke(agents, ex, "sheets.clean.dedupe", {
            "file_ref": ref, "key_columns": ["Customer ID"], "filename": "x"}, chat_level=1)
        assert status == "error" and result.get("denied") is True


class TestReconciling:
    def test_two_lists_match_by_rules_and_near_matches_stay_uncertain(self, agents):
        ex, provider = make()
        a = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        b = upload(provider, "sales.xlsx", sales_xlsx())
        result, status = invoke(agents, ex, "sheets.reconcile.match", {
            "file_ref_a": a, "file_ref_b": b, "filename": "reconciled",
            "keys": [{"a": "Customer ID", "b": "Account", "normalize": "trim"},
                     {"a": "Email", "b": "Contact Email", "normalize": "lowercase"}],
            "compare": [{"a": "Phone", "b": "Mobile"}]})
        assert status == "success", result
        # 00017 (twice in A, once in B: the first matches, the second is
        # only in A), 00019 matches; 00020's email differs → uncertain;
        # 00018 only in A; 00021 and the Total row only in B.
        assert result["matched"] == 2
        assert result["uncertain"] == 1
        assert result["uncertain_pairs"] == [{"row_a": 6, "row_b": 5,
                                              "agree": "Customer ID", "differ": "Email"}]
        assert result["only_in_a"] == 2 and result["only_in_b"] == 2
        assert result["differences"] == 2       # phone spellings differ on both matches
        raw = provider.files["sheets__workbook"][result["file_ref"]]["content"]
        wb = load_workbook(io.BytesIO(raw))
        assert wb.sheetnames == ["Matched", "Only in A", "Only in B", "Uncertain",
                                 "Differences", "Change Log"]
        uncertain = list(wb["Uncertain"].iter_rows(values_only=True))
        assert uncertain[1][:4] == (6, 5, "Customer ID", "Email")
        diffs = list(wb["Differences"].iter_rows(values_only=True))
        assert diffs[1][2] == "Phone" and diffs[1][3] == "+971 50 123 4567"
        log = list(wb["Change Log"].iter_rows(values_only=True))
        assert log[1][1] == "rules"
        assert any(r[1] == "uncertain" and "not merged" in r[2] for r in log)
        job = provider.data["sheets__job"][result["job_ref"]]
        assert job["keys"]["summary"].startswith("2 matched, 2 only in A, 2 only in B, 1 uncertain")

    def test_an_unknown_key_column_is_named(self, agents):
        ex, provider = make()
        a = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        result, status = invoke(agents, ex, "sheets.reconcile.match", {
            "file_ref_a": a, "file_ref_b": a, "filename": "x",
            "keys": [{"a": "Customer ID", "b": "Nope"}]})
        assert status == "error" and "No column 'Nope'" in result["error"]


class TestCombining:
    MARINA = "Date,Merchant,Category,Amount\n2026-09-18,Careem,Transport,38.00\n2026-09-18,Al Seef House,Meals,184.50\n"
    DEIRA = "Date,Merchant,Category,Amount\n2026-09-18,Creekside Suites,Hotel,\"1,240.00\"\n2026-09-18,Pier 9 Supply,Stationery,n/a\n"
    JLT = "Date,Merchant,Amount,Card\n2026-09-18,Northlight,75,00412\n"

    def combine(self, agents, files, **more):
        ex, provider = make()
        refs = [{"file_ref": upload(provider, name, raw.encode()), **extra} for name, raw, extra in files]
        result, status = invoke(agents, ex, "sheets.combine.files",
                                {"files": refs, "filename": "branches", **more})
        return result, status, provider

    def test_several_files_become_one_workbook_with_a_summary(self, agents):
        result, status, provider = self.combine(agents, [
            ("marina.csv", self.MARINA, {"sheet_name": "Marina"}),
            ("deira.csv", self.DEIRA, {}),
            ("jlt.csv", self.JLT, {"sheet_name": "JLT"})], total_column="Amount")
        assert status == "success", result
        assert result["combined"] == 3 and result["left_out"] == 0 and result["rows"] == 5
        assert result["sheets"] == ["Summary", "All", "Marina", "deira", "JLT", "Change Log"]
        assert result["grand_total"] == 1537.5 and result["verified"] is True
        by_source = {f["source"]: f for f in result["files"]}
        assert by_source["Marina"]["total"] == 222.5 and by_source["Marina"]["problem"] == ""
        assert by_source["deira"]["total"] == 1240.0 and by_source["deira"]["not_a_number"] == 1
        assert "not numbers" in by_source["deira"]["problem"]
        assert by_source["JLT"]["problem"] == "missing column(s): Category; extra column(s): Card"

        wb = load_workbook(io.BytesIO(provider.files["sheets__workbook"][result["file_ref"]]["content"]))
        summary = list(wb["Summary"].iter_rows(values_only=True))
        assert summary[0] == ("Source", "File", "Rows", "Total Amount", "Not a number", "Problem")
        assert summary[-1][:4] == ("All files", None, 5, 1537.5)
        everything = list(wb["All"].iter_rows(values_only=True))
        assert everything[0] == ("Source", "Date", "Merchant", "Category", "Amount", "Card")
        assert everything[1] == ("Marina", "2026-09-18", "Careem", "Transport", 38, None)
        assert everything[5] == ("JLT", "2026-09-18", "Northlight", None, 75, "00412")   # zeros kept
        deira = list(wb["deira"].iter_rows(values_only=True))
        assert deira[1][3] == 1240 and deira[2][3] == "n/a"          # a number is a number, the rest is as it came

    def test_a_file_that_cannot_be_read_is_left_out_and_said(self, agents):
        result, status, _ = self.combine(agents, [
            ("marina.csv", self.MARINA, {}), ("empty.csv", "", {}), ("deira.csv", self.DEIRA, {})])
        assert status == "success", result
        assert result["combined"] == 2 and result["left_out"] == 1 and result["grand_total"] is None
        assert [f["problem"] for f in result["files"] if f["filename"] == "empty.csv"] == ["the file is empty"]
        assert "empty" not in result["sheets"]

    def test_two_files_of_one_name_get_two_sheets(self, agents):
        result, status, _ = self.combine(agents, [
            ("receipts.csv", self.MARINA, {}), ("receipts.csv", self.DEIRA, {})])
        assert status == "success", result
        assert result["sheets"] == ["Summary", "All", "receipts", "receipts (2)", "Change Log"]


class TestCalculating:
    def test_aggregate_is_decimal_and_skips_what_is_not_a_number(self, agents):
        ex, provider = make()
        ref = upload(provider, "sales.xlsx", sales_xlsx())
        result, status = invoke(agents, ex, "sheets.calc.aggregate", {
            "file_ref": ref, "measures": [{"column": "Value", "op": "sum"},
                                          {"column": "Orders", "op": "sum"},
                                          {"column": "Company", "op": "count"}]}, chat_level=1)
        assert status == "success", result
        # The Total row's formulas carry no cached value in a fresh file:
        # they are skipped and counted, never guessed.
        assert result["groups"] == [{"sum(Value)": 16080.49, "sum(Orders)": 11.0, "count(Company)": 5}]
        assert result["skipped_cells"] == 2 and result["rows_read"] == 5

        ref = upload(provider, "customers.csv", CUSTOMERS_CSV.encode())
        result, _ = invoke(agents, ex, "sheets.calc.aggregate", {
            "file_ref": ref, "group_by": ["Country"],
            "measures": [{"column": "Name", "op": "count"}, {"column": "Phone", "op": "sum"}]},
            chat_level=1)
        groups = {g["Country"]: g for g in result["groups"]}
        assert groups["UAE"]["count(Name)"] == 3 and groups["uae"]["count(Name)"] == 1
        assert result["skipped_cells"] == 4          # phone numbers are not numbers
        assert groups["UAE"]["sum(Phone)"] is None
