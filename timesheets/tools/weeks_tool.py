"""A week at a glance, locked when it is handed in, and the workbook
that goes with it.

A gap is a weekday, up to as_of, with fewer hours than the day should
hold; days still to come are not gaps. Submitting locks the week's
entries until it is reopened. The workbook is read back before it is
returned.
"""

import base64
import io
from decimal import Decimal

from decentai_sdk.base import ToolBase

from .common import (DAY_LIMIT, DAY_NAMES, as_text, entries, hours_of, iso_date,
                     keys_of, projects, today, week_days, week_of, week_record)


async def summarize(call, week, days, as_of, daily):
    rows = await entries(call, {"week": week})
    known = await projects(call)
    grid = {}
    for entry in rows:
        day = iso_date(entry["date"])
        if day not in days:
            continue
        cells = grid.setdefault(entry["project_code"], [Decimal(0)] * 7)
        cells[days.index(day)] += hours_of(entry["hours"]) or Decimal(0)
    day_totals = [sum((cells[i] for cells in grid.values()), Decimal(0)) for i in range(7)]
    table, billable = [], Decimal(0)
    for code in sorted(grid):
        keys = (known.get(code) or {}).get("keys") or {}
        total = sum(grid[code], Decimal(0))
        if keys.get("billable") == "yes":
            billable += total
        table.append({"project_code": code, "name": str(keys.get("name") or ""),
                      "billable": keys.get("billable") == "yes",
                      "days": [float(h) for h in grid[code]], "total": float(total)})
    gaps = [{"date": day.isoformat(), "logged": float(day_totals[i]),
             "short": float(daily - day_totals[i])}
            for i, day in enumerate(days[:5]) if day <= as_of and day_totals[i] < daily]
    status = keys_of(await week_record(call, week)).get("status") or "open"
    return {"week": week, "from": days[0].isoformat(), "to": days[6].isoformat(),
            "status": status, "dates": [d.isoformat() for d in days],
            "projects": table, "day_totals": [float(t) for t in day_totals],
            "total": float(sum(day_totals, Decimal(0))), "billable_hours": float(billable),
            "gaps": gaps, "entries": len(rows)}


def _week(call):
    """(week, days, as_of, daily, problem)."""
    week = str(call.inputs.get("week") or week_of(today())).strip()
    days = week_days(week)
    as_of = iso_date(call.inputs["as_of"]) if call.inputs.get("as_of") else today()
    daily = hours_of(call.inputs.get("daily_hours", 8))
    if days is None:
        return week, days, as_of, daily, ({"error": "week is an ISO week, like 2026-W37."}, "error")
    if as_of is None:
        return week, days, as_of, daily, ({"error": "as_of must be YYYY-MM-DD."}, "error")
    if daily is None or daily <= 0 or daily > DAY_LIMIT:
        return week, days, as_of, daily, ({"error": "daily_hours must be between 0 and 24."}, "error")
    return week, days, as_of, daily, None


class WeeksTool(ToolBase):
    id = "weeks"

    async def summary(self, call):
        week, days, as_of, daily, problem = _week(call)
        if problem:
            return problem
        return await summarize(call, week, days, as_of, daily), "success"

    async def submit(self, call):
        week, days, as_of, daily, problem = _week(call)
        if problem:
            return problem
        record = await week_record(call, week)
        if keys_of(record).get("status") == "submitted":
            return {"error": f"The week {week} was already submitted on "
                             f"{keys_of(record).get('submitted_on')}."}, "error"
        rows = await entries(call, {"week": week})
        if not rows:
            return {"error": f"Nothing is logged in {week}."}, "error"
        summary = await summarize(call, week, days, as_of, daily)
        for entry in rows:
            if entry["status"] != "submitted":
                await call.resources.update_data("entry", entry["entry_ref"], {"status": "submitted"})
        fields = {"week": week, "status": "submitted", "submitted_on": today().isoformat(),
                  "total_hours": as_text(Decimal(str(summary["total"])))}
        if record:
            await call.resources.update_data("week", record["resource_ref"], fields)
        else:
            await call.resources.create_data("week", fields)
        return {"week": week, "status": "submitted", "total": summary["total"],
                "entries": len(rows), "gaps": summary["gaps"]}, "success"

    async def reopen(self, call):
        week, _, _, _, problem = _week(call)
        if problem:
            return problem
        record = await week_record(call, week)
        if keys_of(record).get("status") != "submitted":
            return {"error": f"The week {week} is not submitted."}, "error"
        await call.resources.update_data("week", record["resource_ref"],
                                         {"status": "open", "submitted_on": ""})
        for entry in await entries(call, {"week": week}):
            await call.resources.update_data("entry", entry["entry_ref"], {"status": "draft"})
        return {"week": week, "status": "open"}, "success"

    async def workbook(self, call):
        from openpyxl import Workbook, load_workbook

        week, days, as_of, daily, problem = _week(call)
        if problem:
            return problem
        summary = await summarize(call, week, days, as_of, daily)
        if not summary["entries"]:
            return {"error": f"Nothing is logged in {week}."}, "error"
        rows = sorted(await entries(call, {"week": week}),
                      key=lambda e: (e["date"], e["project_code"]))

        wb = Workbook()
        sheet = wb.active
        sheet.title = "Timesheet"
        sheet.append([f"Timesheet — week {week}, {days[0]:%d %b} to {days[6]:%d %b %Y} — "
                      f"{summary['status']}"])
        sheet.append([])
        sheet.append(["Project", "Name", "Billable"]
                     + [f"{DAY_NAMES[i]} {d.day:02d}" for i, d in enumerate(days)] + ["Total"])
        for project in summary["projects"]:
            sheet.append([project["project_code"], project["name"],
                          "yes" if project["billable"] else "no"]
                         + project["days"] + [project["total"]])
        sheet.append(["Total", "", ""] + summary["day_totals"] + [summary["total"]])
        detail = wb.create_sheet("Entries")
        detail.append(["Date", "Project", "Hours", "Note", "Source", "Status"])
        for entry in rows:
            detail.append([entry["date"], entry["project_code"], entry["hours"],
                           entry["note"], entry["source"], entry["status"]])
        buffer = io.BytesIO()
        wb.save(buffer)
        raw = buffer.getvalue()
        try:
            verified = load_workbook(io.BytesIO(raw), read_only=True).sheetnames == ["Timesheet", "Entries"]
        except Exception as exc:
            return {"error": f"The workbook did not read back: {exc}"}, "error"
        filename = f"timesheet-{week}.xlsx"
        saved = await call.resources.create_file(
            "workbook", filename, content_base64=base64.b64encode(raw).decode("ascii"))
        return {"file_ref": saved["resource_ref"], "filename": filename,
                "total": summary["total"], "verified": verified}, "success"
