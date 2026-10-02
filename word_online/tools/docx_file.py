"""What is inside a .docx, read from the package itself.

Graph cannot say what a Word document holds, so this does, from the
bytes: paragraphs through python-docx, in reading order with table
cells included, and comments straight from ``word/comments.xml`` with
zipfile and ElementTree — python-docx reads them only in part, and the
format is small and stable. A comment is placed on the paragraph that
holds its ``commentRangeStart`` (or, failing that, its reference mark).

It also makes the one change the agent makes: one paragraph's words
replaced, the paragraph and its first run's formatting kept, and the
package saved whole — every other part (comments, styles, headers)
goes back as it came.
"""

from __future__ import annotations

import copy
import hashlib
import io
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, List, Optional

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

# What a paragraph holds besides words, which replacing its words would
# silently destroy or corrupt: tracked changes, pictures, embedded
# objects, equations. Such a paragraph is edited in Word, not here.
NOT_PLAIN = [qn("w:ins"), qn("w:del"), qn("w:drawing"), qn("w:pict"),
             qn("w:object"), MATH + "oMath"]
# The inline content a paragraph's words live in, all removed when they
# are replaced. Bookmarks and comment range marks are not content and stay.
CONTENT = {qn("w:r"), qn("w:hyperlink"), qn("w:smartTag"), qn("w:fldSimple"), qn("w:sdt")}


class NotEditable(Exception):
    pass


def _is_reference_run(element) -> bool:
    return element.tag == qn("w:r") and element.find(qn("w:commentReference")) is not None


class WordPackage:
    def __init__(self, raw: bytes):
        self.raw = raw
        self.document = docx.Document(io.BytesIO(raw))
        self.paragraphs: List[Paragraph] = []
        self.in_table: List[bool] = []
        self._walk(self.document.element.body, False)

    def _walk(self, container, in_table: bool) -> None:
        parent = self.document._body
        for child in container:
            if child.tag == qn("w:p"):
                self.paragraphs.append(Paragraph(child, parent))
                self.in_table.append(in_table)
            elif child.tag == qn("w:tbl"):
                for row in child.findall(qn("w:tr")):
                    for cell in row.findall(qn("w:tc")):
                        self._walk(cell, True)
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    self._walk(content, in_table)

    # -- reading ---------------------------------------------------------
    @staticmethod
    def style_of(paragraph: Paragraph) -> str:
        try:
            return str(paragraph.style.name) if paragraph.style is not None else "Normal"
        except (KeyError, ValueError):
            return "Normal"

    def row(self, number: int) -> Dict[str, object]:
        paragraph = self.paragraphs[number - 1]
        return {"index": number, "style": self.style_of(paragraph),
                "text": paragraph.text, "in_table": self.in_table[number - 1]}

    def plain(self, number: int) -> bool:
        element = self.paragraphs[number - 1]._p
        return not any(element.find(".//" + tag) is not None for tag in NOT_PLAIN)

    def comments(self) -> List[Dict[str, object]]:
        """Every comment in the package, in file order. ``key`` is a
        fingerprint of author, date and text: Word renumbers comment ids
        when it saves, so the id alone does not say whether a comment is
        one seen before."""
        try:
            xml = zipfile.ZipFile(io.BytesIO(self.raw)).read("word/comments.xml")
        except KeyError:
            return []
        anchors = self._anchors()
        found = []
        for comment in ET.fromstring(xml).iter(W + "comment"):
            comment_id = str(comment.get(W + "id") or "")
            text = "\n".join("".join(t.text or "" for t in p.iter(W + "t"))
                             for p in comment.iter(W + "p")).strip()
            author = str(comment.get(W + "author") or "")
            date = str(comment.get(W + "date") or "")
            number = anchors.get(comment_id, 0)
            found.append({
                "comment_id": comment_id, "author": author, "date": date, "text": text,
                "paragraph": number,
                "paragraph_text": self.paragraphs[number - 1].text if number else "",
                "key": hashlib.sha1(f"{author}\x1f{date}\x1f{text}".encode("utf-8")).hexdigest()[:16],
            })
        return found

    def _anchors(self) -> Dict[str, int]:
        anchors: Dict[str, int] = {}
        for tag in (qn("w:commentReference"), qn("w:commentRangeStart")):
            for number, paragraph in enumerate(self.paragraphs, start=1):
                for mark in paragraph._p.iter(tag):
                    # The range start wins over the reference mark: it is
                    # where the commented words begin.
                    anchors[str(mark.get(qn("w:id")))] = number
        return anchors

    # -- the one change --------------------------------------------------
    def replace_text(self, number: int, new_text: str) -> bytes:
        paragraph = self.paragraphs[number - 1]
        if not self.plain(number):
            raise NotEditable(f"Paragraph {number} holds tracked changes, a picture, an "
                              f"object or an equation; edit it in Word.")
        element = paragraph._p
        first: Optional[object] = next((r for r in element.iter(qn("w:r"))
                                        if not _is_reference_run(r)), None)
        position = None
        for child in list(element):
            if child.tag in CONTENT and not _is_reference_run(child):
                if position is None:
                    position = list(element).index(child)
                element.remove(child)
        run = OxmlElement("w:r")
        if first is not None and first.find(qn("w:rPr")) is not None:
            run.append(copy.deepcopy(first.find(qn("w:rPr"))))
        if position is None:
            # An empty paragraph: after its properties, before any mark.
            position = 1 if element.find(qn("w:pPr")) is not None else 0
        element.insert(position, run)
        Run(run, paragraph).text = new_text
        out = io.BytesIO()
        self.document.save(out)
        return out.getvalue()
