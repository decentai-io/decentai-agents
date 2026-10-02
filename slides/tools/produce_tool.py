"""Producing decks people can open — and reading each one back before it
is returned, so a corrupt file is never handed over as done.

When a template is given, the deck is built on that file: its theme,
fonts and layouts are what a new slide inherits. The template's own
slides are removed first, because a template is a design, not content.
"""

import base64
import io
import re

from decentai_sdk.base import ToolBase

from .reading import kind_of, read_bytes, shape_lines

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_.\- ]+?)\s*\}\}")

TITLE_LAYOUTS = ("title slide", "title")
CONTENT_LAYOUTS = ("title and content", "title, content", "content with caption",
                   "title and body")
BLANK_LAYOUTS = ("title only", "blank")


def _rows(slide):
    rows = slide.get("table_rows") or []
    return [[str(cell) for cell in row] if isinstance(row, (list, tuple)) else [str(row)]
            for row in rows]


def _outline(inputs):
    out = []
    for slide in inputs["slides"]:
        columns = [str(c) for c in slide.get("table_columns") or []]
        rows = _rows(slide)
        if rows and not columns:
            raise ValueError("a table needs table_columns as well as table_rows")
        if rows and any(len(r) != len(columns) for r in rows):
            raise ValueError("every table row must have one cell per column")
        out.append({"title": str(slide.get("title") or ""),
                    "bullets": [str(b) for b in slide.get("bullets") or []],
                    "notes": str(slide.get("notes") or ""),
                    "columns": columns, "rows": rows})
    return out


class ProduceTool(ToolBase):
    id = "produce"

    async def pptx(self, call):
        try:
            outline = _outline(call.inputs)
        except ValueError as exc:
            return {"error": str(exc)}, "error"

        template = None
        template_name = ""
        ref = str(call.inputs.get("template_ref") or "")
        if ref:
            filename, file_type, raw = await read_bytes(call, "source", ref)
            if kind_of(filename, file_type, raw) != "pptx":
                return {"error": "The template is not a PowerPoint (.pptx) "
                                 "file."}, "error"
            template, template_name = raw, filename

        try:
            built = self._build(str(call.inputs["title"]),
                                str(call.inputs.get("subtitle") or ""),
                                outline, template)
        except Exception as exc:
            return {"error": f"The deck could not be built: {exc}"}, "error"

        filename = self._named(str(call.inputs["filename"]))
        result, status = await self._deliver(
            call, filename, built, str(call.inputs.get("purpose") or ""),
            template_name)
        if status == "success":
            result["templated"] = bool(template)
        return result, status

    async def fill_template(self, call):
        from pptx import Presentation

        ref = str(call.inputs["template_ref"])
        filename_in, file_type, raw = await read_bytes(call, "source", ref)
        if kind_of(filename_in, file_type, raw) != "pptx":
            return {"error": "The template is not a PowerPoint (.pptx) "
                             "file."}, "error"
        try:
            presentation = Presentation(io.BytesIO(raw))
        except Exception as exc:
            return {"error": f"The template could not be opened: {exc}"}, "error"

        values = {str(k): str(v) for k, v in (call.inputs.get("values") or {}).items()}
        filled, unfilled = set(), set()

        for slide in presentation.slides:
            for shape in slide.shapes:
                self._fill_shape(shape, values, filled, unfilled)
            try:
                if slide.has_notes_slide:
                    frame = slide.notes_slide.notes_text_frame
                    if frame is not None:
                        self._fill_frame(frame, values, filled, unfilled)
            except Exception:
                pass

        buffer = io.BytesIO()
        presentation.save(buffer)
        filename = self._named(str(call.inputs["filename"]))
        result, status = await self._deliver(
            call, filename, buffer.getvalue(),
            str(call.inputs.get("purpose") or ""), filename_in)
        if status != "success":
            return result, status
        result["filled"] = sorted(filled)
        result["unfilled"] = sorted(unfilled)
        return result, "success"

    # ── building ───────────────────────────────────────────────────────
    @staticmethod
    def _named(filename: str) -> str:
        name = re.sub(r"[^\w.\- ]", "_", filename).strip() or "presentation"
        return name if name.lower().endswith(".pptx") else name + ".pptx"

    @classmethod
    def _build(cls, title, subtitle, outline, template) -> bytes:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        presentation = (Presentation(io.BytesIO(template)) if template
                        else Presentation())
        if template:
            cls._clear(presentation)

        cls._title_slide(presentation, title, subtitle)
        for entry in outline:
            wants_body = bool(entry["bullets"])
            layout = cls._layout(
                presentation,
                CONTENT_LAYOUTS if wants_body else BLANK_LAYOUTS,
                1 if wants_body else 5)
            slide = presentation.slides.add_slide(layout)
            cls._set_title(slide, entry["title"])

            body = cls._body_placeholder(slide)
            if wants_body and body is not None:
                frame = body.text_frame
                frame.clear()
                for index, bullet in enumerate(entry["bullets"]):
                    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    paragraph.text = bullet
                    paragraph.level = 0
            elif wants_body:
                cls._textbox(slide, entry["bullets"], Inches, Pt)
            elif body is not None:
                cls._remove(body)

            if entry["rows"]:
                top = Inches(3.2) if wants_body else Inches(1.8)
                cls._table(slide, entry["columns"], entry["rows"], top,
                           Inches, Pt, presentation)
            if entry["notes"]:
                slide.notes_slide.notes_text_frame.text = entry["notes"]

        buffer = io.BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _clear(presentation) -> None:
        """Every slide the template shipped with, dropped — relationship
        first, then the entry in the slide id list, so the saved package
        keeps no dangling reference."""
        slide_ids = presentation.slides._sldIdLst
        for element in list(slide_ids):
            try:
                presentation.part.drop_rel(element.rId)
            except Exception:
                pass
            slide_ids.remove(element)

    @classmethod
    def _title_slide(cls, presentation, title, subtitle) -> None:
        layout = cls._layout(presentation, TITLE_LAYOUTS, 0)
        slide = presentation.slides.add_slide(layout)
        cls._set_title(slide, title)
        body = cls._body_placeholder(slide)
        if body is not None:
            if subtitle:
                body.text_frame.text = subtitle
            else:
                cls._remove(body)

    @staticmethod
    def _layout(presentation, names, fallback):
        layouts = presentation.slide_layouts
        for wanted in names:
            for layout in layouts:
                if str(layout.name or "").strip().lower() == wanted:
                    return layout
        for wanted in names:
            for layout in layouts:
                if wanted in str(layout.name or "").strip().lower():
                    return layout
        if fallback < len(layouts):
            return layouts[fallback]
        return layouts[0]

    @staticmethod
    def _set_title(slide, text) -> None:
        if not text:
            return
        if slide.shapes.title is not None:
            slide.shapes.title.text = text
            return
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4),
                                       Inches(9.0), Inches(1.0))
        frame = box.text_frame
        frame.text = text
        frame.paragraphs[0].runs[0].font.size = Pt(28)
        frame.paragraphs[0].runs[0].font.bold = True

    @staticmethod
    def _body_placeholder(slide):
        """The placeholder a layout means for content: the first one that
        holds text and is not the title."""
        title = slide.shapes.title
        for placeholder in slide.placeholders:
            if title is not None and placeholder._element is title._element:
                continue
            try:
                if placeholder.has_text_frame:
                    return placeholder
            except Exception:
                continue
        return None

    @staticmethod
    def _remove(shape) -> None:
        element = shape._element
        element.getparent().remove(element)

    @staticmethod
    def _textbox(slide, bullets, Inches, Pt) -> None:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7),
                                       Inches(8.4), Inches(4.5))
        frame = box.text_frame
        frame.word_wrap = True
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            for run in paragraph.runs:
                run.font.size = Pt(18)

    @staticmethod
    def _table(slide, columns, rows, top, Inches, Pt, presentation) -> None:
        left = Inches(0.6)
        width = presentation.slide_width - Inches(1.2)
        height = Inches(0.4) * (len(rows) + 1)
        shape = slide.shapes.add_table(len(rows) + 1, len(columns),
                                       left, top, width, height)
        table = shape.table
        for index, name in enumerate(columns):
            cell = table.cell(0, index)
            cell.text = name
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(12)
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = value
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(12)

    # ── filling ────────────────────────────────────────────────────────
    @classmethod
    def _fill_shape(cls, shape, values, filled, unfilled) -> None:
        try:
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for inner in shape.shapes:
                    cls._fill_shape(inner, values, filled, unfilled)
                return
        except Exception:
            pass
        try:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        cls._fill_frame(cell.text_frame, values, filled, unfilled)
                return
        except Exception:
            pass
        try:
            if shape.has_text_frame:
                cls._fill_frame(shape.text_frame, values, filled, unfilled)
        except Exception:
            pass

    @staticmethod
    def _fill_frame(frame, values, filled, unfilled) -> None:
        for paragraph in frame.paragraphs:
            runs = paragraph.runs
            if not runs:
                continue
            text = "".join(run.text for run in runs)
            names = PLACEHOLDER.findall(text)
            if not names:
                continue
            for name in names:
                (filled if name in values else unfilled).add(name)
            new = PLACEHOLDER.sub(
                lambda match: values.get(match.group(1), match.group(0)), text)
            if new == text:
                continue
            # The replacement lands in the first run; the paragraph keeps
            # its own style, character formatting within it may not.
            runs[0].text = new
            for run in runs[1:]:
                run.text = ""

    # ── delivering ─────────────────────────────────────────────────────
    async def _deliver(self, call, filename, raw, purpose, template_name):
        saved = await call.resources.create_file(
            "deck", filename,
            content_base64=base64.b64encode(raw).decode("ascii"))
        file_ref = saved["resource_ref"]
        try:
            count = self._verify(raw)
        except Exception as exc:
            return {"error": f"The produced deck did not read back: {exc}"}, "error"
        record = await call.resources.create_data("output", {
            "filename": filename, "file_ref": file_ref, "slides": count,
            "purpose": purpose[:200], "template": template_name[:200]})
        return {"file_ref": file_ref, "filename": filename,
                "output_ref": record["resource_ref"], "verified": True,
                "slides": count}, "success"

    @staticmethod
    def _verify(raw: bytes) -> int:
        """Open what was written with the same library that wrote it, and
        insist there are slides with words on them."""
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(raw))
        count = len(presentation.slides)
        if count == 0:
            raise ValueError("the deck came back with no slides")
        words = 0
        for slide in presentation.slides:
            for shape in slide.shapes:
                words += sum(len(line) for line in shape_lines(shape))
        if words == 0:
            raise ValueError("no text came back from the deck")
        return count
