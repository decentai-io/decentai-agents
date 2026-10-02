"""Turning a stored deck into slides of text, and saying which slides
had nothing to read.

A slide whose content is a picture, a chart or a diagram says nothing to
a text reader. It is ``empty`` here, by number. It is never dropped: the
caller sees that slide 4 of 30 carries no words, which is different from
slide 4 not existing.
"""

from __future__ import annotations

import base64
import io
from typing import List, Tuple

MAX_SLIDES = 300

PPTX_MIME = ("application/vnd.openxmlformats-officedocument"
             ".presentationml.presentation")


class Slide:
    """One slide, read: what it says and what the speaker was told."""

    def __init__(self, number: int, title: str, text: str, notes: str,
                 layout: str = ""):
        self.number = number
        self.title = title
        self.text = text
        self.notes = notes
        self.layout = layout

    @property
    def empty(self) -> bool:
        return not (self.title or self.text or self.notes)

    def characters(self, with_notes: bool = True) -> int:
        total = len(self.title) + len(self.text)
        return total + len(self.notes) if with_notes else total


class Deck:
    """One file, read: its kind, its slides, and what failed."""

    def __init__(self, file_ref: str, filename: str, kind: str,
                 slides: List[Slide], problem: str = ""):
        self.file_ref = file_ref
        self.filename = filename
        self.kind = kind
        self.slides = slides
        self.problem = problem

    @property
    def empty_slides(self) -> List[int]:
        return [s.number for s in self.slides if s.empty]

    @property
    def characters(self) -> int:
        return sum(s.characters() for s in self.slides)

    @property
    def has_notes(self) -> bool:
        return any(s.notes for s in self.slides)

    @property
    def layouts(self) -> List[str]:
        seen: List[str] = []
        for slide in self.slides:
            if slide.layout and slide.layout not in seen:
                seen.append(slide.layout)
        return seen


def kind_of(filename: str, file_type: str, raw: bytes) -> str:
    name = (filename or "").lower()
    mime = (file_type or "").lower()
    if raw[:2] == b"PK" and (name.endswith(".pptx") or mime == PPTX_MIME
                             or b"ppt/" in raw[:4000]):
        return "pptx"
    return "unsupported"


async def read_bytes(call, resource_id: str, file_ref: str) -> Tuple[str, str, bytes]:
    """The stored file as (filename, declared type, bytes)."""
    record = await call.resources.read_file(resource_id, file_ref)
    filename = str(record.get("filename") or "")
    file_type = str(record.get("file_type") or record.get("mime_type") or "")
    if record.get("content_base64"):
        raw = base64.b64decode(record["content_base64"])
    else:
        raw = str(record.get("content") or "").encode("utf-8")
    return filename, file_type, raw


async def load(call, resource_id: str, file_ref: str) -> Deck:
    filename, file_type, raw = await read_bytes(call, resource_id, file_ref)
    kind = kind_of(filename, file_type, raw)
    if kind != "pptx":
        return Deck(file_ref, filename, kind, [],
                    problem="Not a PowerPoint (.pptx) file. Legacy .ppt "
                            "files are not read.")
    slides, problem = _pptx(raw)
    return Deck(file_ref, filename, kind, slides, problem)


def _pptx(raw: bytes) -> Tuple[List[Slide], str]:
    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(raw))
    except Exception as exc:                      # a corrupt or foreign package
        return [], f"The deck could not be opened: {exc}"

    count = len(presentation.slides)
    if count > MAX_SLIDES:
        return [], f"The deck has {count} slides; at most {MAX_SLIDES} are read."

    slides: List[Slide] = []
    for number, slide in enumerate(presentation.slides, 1):
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None
        title = _clean(title_shape.text) if title_shape is not None else ""

        skip = id(title_shape) if title_shape is not None else 0
        lines: List[str] = []
        for shape in slide.shapes:
            if id(shape) == skip:
                continue
            lines.extend(shape_lines(shape))

        notes = ""
        try:
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                notes = _clean(frame.text) if frame is not None else ""
        except Exception:
            notes = ""

        try:
            layout = str(slide.slide_layout.name or "")
        except Exception:
            layout = ""

        slides.append(Slide(number, title, "\n".join(lines), notes, layout))
    return slides, ""


def shape_lines(shape) -> List[str]:
    """Every line of text a shape carries: its paragraphs, its table's
    rows, and the shapes inside it when it is a group."""
    lines: List[str] = []
    if _is_group(shape):
        for inner in shape.shapes:
            lines.extend(shape_lines(inner))
        return lines
    try:
        if shape.has_table:
            for row in shape.table.rows:
                cells = [_clean(cell.text) for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
            return lines
    except Exception:
        pass
    try:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = _clean("".join(run.text for run in paragraph.runs))
                if text:
                    lines.append(text)
    except Exception:
        pass
    return lines


def _is_group(shape) -> bool:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        return shape.shape_type == MSO_SHAPE_TYPE.GROUP
    except Exception:
        return False


def _clean(text: str) -> str:
    """A vertical tab is what PowerPoint writes for a soft line break;
    it becomes a newline, and trailing space goes."""
    return str(text or "").replace("\v", "\n").strip()
