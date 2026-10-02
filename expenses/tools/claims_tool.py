import base64
import io
from collections import defaultdict
from decimal import Decimal

from decentai_sdk.base import ToolBase

from .common import (iso_date, keys_of, money, normalize_merchant, out, receipt_row,
                     rules_of, today)


class ClaimsTool(ToolBase):
    id = "claims"

    async def list(self, call):
        wanted = str(call.inputs.get("status") or "")
        rows = await call.resources.list_data("claim", {"status": wanted} if wanted else None)
        return {"claims": [{
            "claim_ref": r["resource_ref"],
            **{k: str(keys_of(r).get(k) or "")
               for k in ("title", "claimant", "status", "currency", "note")},
        } for r in rows]}, "success"

    async def create(self, call):
        record = await call.resources.create_data("claim", {
            "title": str(call.inputs["title"]).strip(),
            "claimant": str(call.inputs["claimant"]).strip(), "status": "draft",
            "currency": str(call.inputs.get("currency") or "USD").upper()})
        return {"claim_ref": record["resource_ref"], "status": "draft"}, "success"

    async def _receipts(self, call, claim_ref):
        rows = [receipt_row(r) for r in await call.resources.list_data("receipt", {"claim_ref": claim_ref})]
        rows.sort(key=lambda r: (r["date"] or "9999", r["filename"]))
        return rows

    async def get(self, call):
        ref = str(call.inputs["claim_ref"])
        claim = keys_of(await call.resources.read_data("claim", ref))
        receipts = await self._receipts(call, ref)
        totals = defaultdict(lambda: [Decimal("0"), 0])
        for row in receipts:
            if row["status"] == "duplicate" or row.get("amount") is None:
                continue
            totals[row["currency"] or "?"][0] += money(row["amount"]) or Decimal("0")
            totals[row["currency"] or "?"][1] += 1
        return {
            "claim_ref": ref, "title": str(claim.get("title") or ""),
            "claimant": str(claim.get("claimant") or ""), "status": str(claim.get("status") or "draft"),
            "currency": str(claim.get("currency") or "USD"),
            "decided_by": str(claim.get("decided_by") or ""), "decided_on": str(claim.get("decided_on") or ""),
            "needs_attention": sum(1 for r in receipts if r["status"] in ("needs_review", "manual", "duplicate")),
            "totals": [{"currency": c, "amount": out(t[0]), "receipts": t[1]}
                       for c, t in sorted(totals.items())],
            "receipts": [{k: r[k] for k in r if k != "claim_ref"} for r in receipts],
        }, "success"

    async def _exceptions(self, call, claim_ref, as_of):
        """Every reason a reviewer would stop on a receipt, from the
        rules in force and the receipts as recorded. Deterministic."""
        claim = keys_of(await call.resources.read_data("claim", claim_ref))
        claim_currency = str(claim.get("currency") or "USD")
        receipts = await self._receipts(call, claim_ref)
        rules = rules_of(await call.resources.list_data("rule", {}))
        exceptions = []

        def add(row, kind, detail):
            exceptions.append({"receipt_ref": row["receipt_ref"], "filename": row["filename"],
                               "kind": kind, "detail": detail})

        # Duplicates across everything recorded, not only this claim.
        everything = [receipt_row(r) for r in await call.resources.list_data("receipt", {})]
        seen = {}
        for row in sorted(everything, key=lambda r: r["receipt_ref"]):
            if row["status"] == "duplicate" or row.get("amount") is None or not row["date"]:
                continue
            key = (normalize_merchant(row["merchant"]), row["date"], row["amount"])
            if all(key) and key in seen and seen[key] != row["receipt_ref"]:
                if row["claim_ref"] == claim_ref:
                    add(row, "duplicate", f"same merchant, date and amount as {seen[key]}")
                    await call.resources.update_data("receipt", row["receipt_ref"],
                                                     {"status": "duplicate", "duplicate_of": seen[key]})
            else:
                seen.setdefault(key, row["receipt_ref"])

        meals_by_day = defaultdict(Decimal)
        for row in receipts:
            if row["status"] == "duplicate":
                add(row, "duplicate", f"same merchant, date and amount as {row['duplicate_of'] or 'another receipt'}")
                continue
            if row["status"] == "manual":
                add(row, "manual", "a photo: details must be entered and confirmed")
                continue
            if row["status"] == "needs_review":
                add(row, "needs_review", row["note"] or f"assumed: {row['assumed']}")
            if row["currency"] and row["currency"] != claim_currency:
                add(row, "currency", f"{row['currency']} receipt in a {claim_currency} claim; not converted")
            amount = money(row.get("amount")) if row.get("amount") is not None else None
            if amount is None:
                continue
            if row["category"] == "meals" and row["date"] and row["currency"] == claim_currency:
                meals_by_day[row["date"]] += amount
            if row["category"] == "accommodation" and "hotel_max_per_night" in rules:
                cap = money(rules["hotel_max_per_night"][0])
                if cap is not None and amount > cap:
                    add(row, "policy", f"accommodation {out(amount)} exceeds hotel_max_per_night {out(cap)}")
            if row["category"] == "transport" and "taxi_max_per_trip" in rules:
                cap = money(rules["taxi_max_per_trip"][0])
                if cap is not None and amount > cap:
                    add(row, "policy", f"transport {out(amount)} exceeds taxi_max_per_trip {out(cap)}")
            if "claim_within_days" in rules and row["date"]:
                days = rules["claim_within_days"][0]
                receipt_date = iso_date(row["date"])
                if days.isdigit() and receipt_date and (as_of - receipt_date).days > int(days):
                    add(row, "policy", f"receipt dated {row['date']} is older than claim_within_days {days}")
        if "meal_max_per_day" in rules:
            cap = money(rules["meal_max_per_day"][0])
            if cap is not None:
                for day, total in sorted(meals_by_day.items()):
                    if total > cap:
                        for row in receipts:
                            if row["category"] == "meals" and row["date"] == day and row["status"] != "duplicate":
                                add(row, "policy", f"meals on {day} total {out(total)}, over meal_max_per_day {out(cap)}")
                                break
        applied = sorted(f"{name} = {value}" for name, (value, _) in rules.items())
        return receipts, exceptions, applied

    async def check(self, call):
        ref = str(call.inputs["claim_ref"])
        as_of = iso_date(call.inputs["as_of"]) if call.inputs.get("as_of") else today()
        if as_of is None:
            return {"error": "as_of must be YYYY-MM-DD."}, "error"
        receipts, exceptions, applied = await self._exceptions(call, ref, as_of)
        flagged = {e["receipt_ref"] for e in exceptions}
        return {"claim_ref": ref, "receipts": len(receipts),
                "ok": sum(1 for r in receipts if r["receipt_ref"] not in flagged),
                "exceptions": exceptions, "rules_applied": applied}, "success"

    async def workbook(self, call):
        from openpyxl import Workbook, load_workbook

        ref = str(call.inputs["claim_ref"])
        as_of = iso_date(call.inputs["as_of"]) if call.inputs.get("as_of") else today()
        if as_of is None:
            return {"error": "as_of must be YYYY-MM-DD."}, "error"
        claim = keys_of(await call.resources.read_data("claim", ref))
        receipts, exceptions, applied = await self._exceptions(call, ref, as_of)
        receipts = await self._receipts(call, ref)         # statuses after the check

        wb = Workbook()
        ws = wb.active
        ws.title = "Receipts"
        ws.append([f"Expense claim — {claim.get('title')} — {claim.get('claimant')} — status {claim.get('status')}"])
        ws.append([f"Prepared {as_of.isoformat()}; amounts as on the receipts, no conversion"])
        ws.append([])
        ws.append(["Date", "Merchant", "Category", "Currency", "Amount", "Status", "Verified",
                   "Assumed", "Duplicate of", "Note", "File", "Receipt ref"])
        for row in receipts:
            ws.append([row["date"], row["merchant"], row["category"], row["currency"],
                       row.get("amount"), row["status"], row["verified"], row["assumed"],
                       row["duplicate_of"], row["note"], row["filename"], row["receipt_ref"]])
        summary = wb.create_sheet("Summary")
        summary.append(["Category", "Currency", "Receipts", "Total"])
        by = defaultdict(lambda: [0, Decimal("0")])
        for row in receipts:
            if row["status"] == "duplicate" or row.get("amount") is None:
                continue
            by[(row["category"], row["currency"])][0] += 1
            by[(row["category"], row["currency"])][1] += money(row["amount"]) or Decimal("0")
        for (category, currency), (count, total) in sorted(by.items()):
            summary.append([category, currency, count, out(total)])
        exc = wb.create_sheet("Exceptions")
        exc.append(["Receipt ref", "File", "Kind", "Detail"])
        for e in exceptions:
            exc.append([e["receipt_ref"], e["filename"], e["kind"], e["detail"]])
        rules = wb.create_sheet("Rules applied")
        rules.append(["Rule"])
        for line in applied:
            rules.append([line])
        buffer = io.BytesIO()
        wb.save(buffer)
        raw = buffer.getvalue()
        filename = f"claim-{str(claim.get('claimant') or 'claim').replace(' ', '-')}-{as_of.isoformat()}.xlsx"
        saved = await call.resources.create_file(
            "workbook", filename, content_base64=base64.b64encode(raw).decode("ascii"))
        try:
            verified = "Receipts" in load_workbook(io.BytesIO(raw), read_only=True).sheetnames
        except Exception as exc_:
            return {"error": f"The workbook did not read back: {exc_}"}, "error"
        return {"file_ref": saved["resource_ref"], "filename": filename,
                "receipts": len(receipts), "exceptions": len(exceptions),
                "verified": verified}, "success"

    async def set_status(self, call):
        ref = str(call.inputs["claim_ref"])
        claim = keys_of(await call.resources.read_data("claim", ref))
        current = str(claim.get("status") or "draft")
        wanted = str(call.inputs["status"])
        by = str(call.inputs.get("by") or "").strip()
        on = iso_date(call.inputs["on_date"]) if call.inputs.get("on_date") else today()
        if on is None:
            return {"error": "on_date must be YYYY-MM-DD."}, "error"
        allowed = {"draft": {"submitted"}, "submitted": {"approved", "rejected"},
                   "approved": {"paid"}, "rejected": set(), "paid": set()}
        if wanted not in allowed[current]:
            return {"error": f"A {current} claim cannot become {wanted}."}, "error"
        if wanted == "submitted":
            receipts = await self._receipts(call, ref)
            pending = [r["filename"] or r["receipt_ref"] for r in receipts
                       if r["status"] in ("needs_review", "manual")]
            if pending:
                return {"error": "Receipts still need review or entry: " + ", ".join(pending)}, "error"
            if not receipts:
                return {"error": "The claim has no receipts."}, "error"
        elif not by:
            return {"error": f"{wanted} is a reviewer's decision: name who decided in \"by\"."}, "error"
        changes = {"status": wanted, "note": str(call.inputs.get("note") or "")}
        if wanted != "submitted":
            changes.update({"decided_by": by, "decided_on": on.isoformat()})
        await call.resources.update_data("claim", ref, changes)
        return {"claim_ref": ref, "status": wanted, "decided_by": changes.get("decided_by", ""),
                "decided_on": changes.get("decided_on", "")}, "success"
