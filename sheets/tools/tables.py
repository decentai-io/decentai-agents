"""A spreadsheet as tables of cells that remember what they were.

A cell keeps its original value (a number, a date, a string, a formula)
and its text as a person sees it. Identifiers with leading zeros stay
text; dates stay dates; a formula stays a formula, with the cached
value the file carried. Comparisons happen on normalized text;
arithmetic happens on Decimal; writing puts the original value back.
"""

from __future__ import annotations

import base64
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 50000
_LEADING_ZERO = re.compile(r"^0\d+$")
_NUMBER = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$")


class Cell:
    __slots__ = ("value", "formula", "cached")

    def __init__(self, value: Any, formula: str = "", cached: Any = None):
        self.value = value          # what the file holds (number/date/str/None)
        self.formula = formula      # "=SUM(A2:A9)" when the cell is a formula
        self.cached = cached        # the formula's last computed value, if any

    @property
    def text(self) -> str:
        shown = self.cached if self.formula else self.value
        if shown is None:
            return ""
        if isinstance(shown, datetime):
            return shown.date().isoformat() if shown.time() == datetime.min.time() else shown.isoformat(sep=" ")
        if isinstance(shown, date):
            return shown.isoformat()
        if isinstance(shown, bool):
            return "TRUE" if shown else "FALSE"
        if isinstance(shown, float) and shown.is_integer():
            return str(int(shown))
        return str(shown)

    @property
    def blank(self) -> bool:
        return not self.formula and (self.value is None or str(self.value).strip() == "")

    def number(self) -> Optional[Decimal]:
        shown = self.cached if self.formula else self.value
        if isinstance(shown, bool) or shown is None:
            return None
        if isinstance(shown, (int, float)):
            return Decimal(str(shown))
        text = str(shown).strip().replace(",", "")
        if not text or not _NUMBER.match(text.replace(",", "")):
            return None
        try:
            return Decimal(text)
        except InvalidOperation:
            return None


class Sheet:
    def __init__(self, name: str, columns: List[str], rows: List[List[Cell]]):
        self.name = name
        self.columns = columns
        self.rows = rows            # index 0 = sheet row 2 (under the header)

    def column(self, name: str) -> int:
        wanted = name.strip().lower()
        for index, column in enumerate(self.columns):
            if column.strip().lower() == wanted:
                return index
        raise KeyError(f"No column '{name}' in sheet '{self.name}' "
                       f"(columns: {', '.join(self.columns)})")

    def cell(self, row: List[Cell], name: str) -> Cell:
        index = self.column(name)
        return row[index] if index < len(row) else Cell(None)

    def row_number(self, index: int) -> int:
        return index + 2

    def profile(self) -> List[Dict[str, Any]]:
        out = []
        for index, name in enumerate(self.columns):
            cells = [r[index] if index < len(r) else Cell(None) for r in self.rows]
            kinds = set()
            blanks = formulas = 0
            zeros = False
            for cell in cells:
                if cell.formula:
                    formulas += 1
                if cell.blank:
                    blanks += 1
                    continue
                shown = cell.cached if cell.formula else cell.value
                if shown is None:
                    continue            # a formula never computed says nothing
                if isinstance(shown, (datetime, date)):
                    kinds.add("date")
                elif isinstance(shown, bool):
                    kinds.add("text")
                elif isinstance(shown, (int, float)):
                    kinds.add("number")
                else:
                    text = str(shown).strip()
                    if _LEADING_ZERO.match(text):
                        zeros = True
                        kinds.add("text")
                    elif cell.number() is not None:
                        kinds.add("number")
                    else:
                        kinds.add("text")
            kind = "empty" if not kinds else (kinds.pop() if len(kinds) == 1 else "mixed")
            out.append({"name": name, "type": kind, "blanks": blanks,
                        "formulas": formulas, "leading_zeros": zeros})
        return out


class Workbook:
    def __init__(self, file_ref: str, filename: str, kind: str, sheets: List[Sheet],
                 problem: str = ""):
        self.file_ref = file_ref
        self.filename = filename
        self.kind = kind
        self.sheets = sheets
        self.problem = problem

    def sheet(self, name: str = "") -> Sheet:
        if not self.sheets:
            raise KeyError("The file has no sheets.")
        if not name:
            return self.sheets[0]
        for sheet in self.sheets:
            if sheet.name.strip().lower() == name.strip().lower():
                return sheet
        raise KeyError(f"No sheet '{name}' (sheets: "
                       f"{', '.join(s.name for s in self.sheets)})")


def normalize(text: str, how: str) -> str:
    text = str(text or "")
    if how == "none":
        return text
    if how == "digits":
        return re.sub(r"\D", "", text)
    text = " ".join(text.split())
    return text.lower() if how == "lowercase" else text


async def load(call, resource_id: str, file_ref: str) -> Workbook:
    record = await call.resources.read_file(resource_id, file_ref)
    filename = str(record.get("filename") or "")
    file_type = str(record.get("file_type") or record.get("mime_type") or "")
    raw = base64.b64decode(record["content_base64"]) if record.get("content_base64") \
        else str(record.get("content") or "").encode("utf-8")
    name = filename.lower()
    if raw[:2] == b"PK" or name.endswith(".xlsx") or file_type == XLSX_MIME:
        return _xlsx(file_ref, filename, raw)
    if name.endswith(".csv") or file_type == "text/csv" or _looks_like_text(raw):
        return _csv(file_ref, filename, raw)
    return Workbook(file_ref, filename, "unsupported", [],
                    "Not a CSV or an Excel .xlsx file.")


def _looks_like_text(raw: bytes) -> bool:
    try:
        raw[:4096].decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def _csv(file_ref: str, filename: str, raw: bytes) -> Workbook:
    text = raw.decode("utf-8-sig", "replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = list(csv.reader(io.StringIO(text), dialect))
    if not reader:
        return Workbook(file_ref, filename, "csv", [Sheet("Sheet1", [], [])])
    columns = [str(c).strip() for c in reader[0]]
    rows = [[Cell(value if value != "" else None) for value in row] for row in reader[1:MAX_ROWS + 1]]
    return Workbook(file_ref, filename, "csv", [Sheet("Sheet1", columns, rows)])


def _xlsx(file_ref: str, filename: str, raw: bytes) -> Workbook:
    from openpyxl import load_workbook

    try:
        formulas = load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        values = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        return Workbook(file_ref, filename, "xlsx", [],
                        f"The workbook could not be opened: {exc}")
    sheets = []
    for name in formulas.sheetnames:
        ws_f, ws_v = formulas[name], values[name]
        rows_f = list(ws_f.iter_rows(values_only=True, max_row=MAX_ROWS + 1))
        rows_v = list(ws_v.iter_rows(values_only=True, max_row=MAX_ROWS + 1))
        if not rows_f:
            sheets.append(Sheet(name, [], []))
            continue
        columns = [str(c).strip() if c is not None else f"Column{i + 1}"
                   for i, c in enumerate(rows_f[0])]
        rows = []
        for row_f, row_v in zip(rows_f[1:], rows_v[1:]):
            cells = []
            for value_f, value_v in zip(row_f, row_v):
                if isinstance(value_f, str) and value_f.startswith("="):
                    cells.append(Cell(None, value_f, value_v))
                else:
                    cells.append(Cell(value_f))
            rows.append(cells)
        while rows and all(c.blank for c in rows[-1]):
            rows.pop()
        sheets.append(Sheet(name, columns, rows))
    return Workbook(file_ref, filename, "xlsx", sheets)


def row_dict(sheet: Sheet, row: List[Cell], columns: Optional[List[str]] = None,
             formulas: bool = False) -> Dict[str, str]:
    """The row keyed by column name, narrowed to ``columns`` when given.

    The narrowing matches a name the way ``Sheet.column`` does — trimmed
    and without regard to case — because everywhere else in this agent
    does, and because the header a person reads is not a key they typed.
    It used to compare the caller's spelling to the header exactly: ask
    for "Title" where the header says "title" and every column was
    dropped, so each row came back as {} — the right NUMBER of rows,
    every one of them empty, reported as success. The guard that checks
    the names are real is itself case-insensitive, so it waved the
    mismatch through on its way here."""
    wanted = {str(c).strip().lower() for c in columns} if columns else set()
    out = {}
    for index, name in enumerate(sheet.columns):
        if wanted and name.strip().lower() not in wanted:
            continue
        cell = row[index] if index < len(row) else Cell(None)
        out[name] = cell.formula if (formulas and cell.formula) else cell.text
    return out


def key_of(sheet: Sheet, row: List[Cell], columns: List[str], how: str) -> Tuple[str, ...]:
    return tuple(normalize(sheet.cell(row, c).text, how) for c in columns)
